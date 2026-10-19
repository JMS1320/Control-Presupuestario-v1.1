"""
📜 PASO 1 del histórico MSA (A-DEC-1009 / A-DAT-68): cargar el extracto jul-2025 → ene-2026 SIN CLASIFICAR.

    python scripts/historico_paso1_extracto.py            → arma el Excel con las filas EXACTAS que se cargarían (no escribe)
    python scripts/historico_paso1_extracto.py --aplicar  → además las inserta en public.msa_galicia (SÓLO con OK del usuario)

Fuente: «Extracto MSA - Control Presupuestario.xlsx» (la planilla del usuario, verificada el 15/10: mismos 860
movimientos jul–ene que la anterior y cadena de saldos que cierra). Cada movimiento entra:
  · estado 'historico', SIN categ, SIN contable/interno (no entra en ningún número hasta revisar su lote);
  · con todas las columnas del banco tal cual, su saldo, y el `control` de la cadena (tiene que dar 0);
  · `orden` NEGATIVO (−860 … −1, del más viejo al más nuevo): la app ordena por fecha y orden, y el importador
    sigue numerando desde el máximo (1.032) sin enterarse;
  · su Detalle en `detalle`, y su CATEG / Contable / Interno en `nota_operador` (dato suyo, hasta revisar el lote).
Deshacer: la app no tiene NINGUNA fila antes del 02/02/2026 → `delete … where estado='historico' and fecha < '2026-02-01'`
(y la lista de ids queda guardada en el Excel).
"""
import openpyxl, datetime, json, sys, urllib.request, urllib.parse
from openpyxl.styles import PatternFill, Font

APLICAR = '--aplicar' in sys.argv
RAIZ = 'D:/Users/josem/Documents/Jose/Automatizarr/Claude/Control-Presupuestario-v1.1/'
DIR = RAIZ + '- Comunicacion JMS Claude - Archivos/Balance/Excels trabajando con claude/'
FUENTE = DIR + 'Extracto MSA - Control Presupuestario.xlsx'
SALIDA = DIR + 'Paso 1 - extracto historico a cargar (jul25-ene26).xlsx'
INICIO, FIN = datetime.datetime(2025, 7, 1), datetime.datetime(2026, 1, 31)
SALDO_INICIO = 832605.05            # su planilla: «saldo de inicio de ej 2025/26 banco galicia cta cte»

env = {}
for l in open(RAIZ + '.env.local', encoding='utf-8'):
    if '=' in l and not l.strip().startswith('#'):
        k, v = l.split('=', 1); env[k.strip()] = v.strip().strip('"').strip("'")
H_API = {'apikey': env['SUPABASE_SERVICE_ROLE_KEY'], 'Authorization': 'Bearer ' + env['SUPABASE_SERVICE_ROLE_KEY'], 'Content-Type': 'application/json'}
def rest(tabla, params):
    url = env['NEXT_PUBLIC_SUPABASE_URL'] + '/rest/v1/' + tabla + '?' + urllib.parse.urlencode(params, safe='(),.*:')
    return json.loads(urllib.request.urlopen(urllib.request.Request(url, headers=H_API)).read())

r2 = lambda x: round(float(x), 2) if isinstance(x, (int, float)) else 0.0
def texto(v):
    """Texto del banco tal cual; un número entero que Excel guardó como número vuelve a texto sin «.0»."""
    if v is None or v == '': return None
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v).strip() or None

# ── 1. La planilla ──────────────────────────────────────────────────────────────────────────────────────
ws = openpyxl.load_workbook(FUENTE, data_only=True).worksheets[0]
HR = 4
H = {str(c.value).strip(): c.column for c in ws[HR] if c.value}
filas = []
for i in range(HR + 1, ws.max_row + 1):
    f = ws.cell(i, 1).value
    if isinstance(f, datetime.datetime) and INICIO <= f <= FIN:
        filas.append(i)
# La planilla va del más nuevo al más viejo (y dentro del día también): se invierte para cargar en orden cronológico.
filas.sort(reverse=True)
g = lambda i, n: ws.cell(i, H[n]).value

registros, saltos = [], []
saldo_prev = SALDO_INICIO
for k, i in enumerate(filas):
    deb, cre, saldo = r2(g(i, 'Débitos')), r2(g(i, 'Créditos')), r2(g(i, 'Saldo'))
    control = round(saldo_prev + cre - deb - saldo, 2)
    if abs(control) > 0.005: saltos.append((i, control))
    nota = ' · '.join(x for x in [
        f"CATEG {texto(g(i, 'CATEG'))}" if texto(g(i, 'CATEG')) else None,
        f"Contable {texto(g(i, 'Contable'))}" if texto(g(i, 'Contable')) else None,
        f"Interno {texto(g(i, 'Interno'))}" if texto(g(i, 'Interno')) else None] if x)
    registros.append({
        'fecha': g(i, 'Fecha').strftime('%Y-%m-%d'),
        'descripcion': texto(g(i, 'Descripción')), 'origen': texto(g(i, 'Origen')) or '',
        'debitos': deb, 'creditos': cre,
        'grupo_de_conceptos': texto(g(i, 'Grupo de Conceptos')), 'concepto': texto(g(i, 'Concepto')),
        'numero_de_terminal': texto(g(i, 'Número de Terminal')), 'observaciones_cliente': texto(g(i, 'Observaciones Cliente')),
        'numero_de_comprobante': texto(g(i, 'Número de Comprobante')),
        'leyendas_adicionales_1': texto(g(i, 'Leyendas Adicionales1')), 'leyendas_adicionales_2': texto(g(i, 'Leyendas Adicionales2')),
        'leyendas_adicionales_3': texto(g(i, 'Leyendas Adicionales3')), 'leyendas_adicionales_4': texto(g(i, 'Leyendas Adicionales4')),
        'tipo_de_movimiento': texto(g(i, 'Tipo de Movimiento')),
        'saldo': saldo, 'control': control,
        'categ': None, 'detalle': texto(g(i, 'Detalle')), 'contable': None, 'interno': None, 'centro_de_costo': None,
        'cuenta': 'MSA Galicia', 'orden': -(len(filas) - k), 'estado': 'historico', 'revisado': False,
        'nota_operador': ('📜 Histórico jul25-ene26 — tu planilla: ' + nota) if nota else '📜 Histórico jul25-ene26',
        '_fila': i,
    })
    saldo_prev = saldo

# ── 2. Controles ────────────────────────────────────────────────────────────────────────────────────────
app = rest('msa_galicia', {'select': 'fecha,saldo,orden,debitos', 'order': 'fecha.asc,orden.asc', 'limit': '1'})[0]
antes = rest('msa_galicia', {'select': 'id', 'fecha': 'lt.2026-02-01'})
ultimo = registros[-1]['saldo']
empalme = round(ultimo - float(app['debitos']) - float(app['saldo']), 2)   # el primero de la app es un débito de comisión
tot_deb = round(sum(r['debitos'] for r in registros), 2); tot_cre = round(sum(r['creditos'] for r in registros), 2)
cierre_calc = round(SALDO_INICIO - tot_deb + tot_cre, 2)
controles = [
    ('Movimientos a cargar (jul-2025 → ene-2026)', len(registros), 860, len(registros) == 860),
    ('Saltos en la cadena de saldos (fila por fila)', len(saltos), 0, not saltos),
    ('Saldo de inicio del ejercicio (30/06/2025, tu planilla)', SALDO_INICIO, SALDO_INICIO, True),
    ('Saldo inicio − débitos + créditos = saldo del último movimiento (30/01/2026)', cierre_calc, ultimo, abs(cierre_calc - ultimo) < 0.01),
    (f"Empalme con la app: saldo 30/01 − 1er débito de la app ({app['debitos']}) = saldo del 02/02/2026 en la app ({app['saldo']})", empalme, 0, abs(empalme) < 0.01),
    ('Filas que la app YA tiene antes del 01/02/2026 (tiene que ser 0: no hay nada que se pise)', len(antes), 0, len(antes) == 0),
    ('Orden: del −860 al −1, sin repetir (el importador sigue en el 1.033)', f"{registros[0]['orden']} … {registros[-1]['orden']}", '−860 … −1', len({r['orden'] for r in registros}) == 860),
]

# ── 3. Excel ────────────────────────────────────────────────────────────────────────────────────────────
AZUL = PatternFill('solid', fgColor='DDEBF7'); VERDE = PatternFill('solid', fgColor='E2EFDA'); ROJO = PatternFill('solid', fgColor='F8CBAD'); B = Font(bold=True)
wb = openpyxl.Workbook(); s0 = wb.active; s0.title = '00 Resumen y controles'
s0.append(['PASO 1 — el extracto de MSA jul-2025 → ene-2026 que se cargaría en la app como HISTÓRICO, SIN CLASIFICAR. NO está cargado: es para que lo revises.'])
s0.append(['Cada movimiento entra con estado «historico», sin categoría y sin códigos contable/interno: no aparece en el dashboard ni en los papeles hasta que revisemos su lote. Sólo aporta la cadena de saldos y el saldo de inicio. Tu CATEG/Contable/Interno van en la nota; tu Detalle, en Detalle.'])
s0.append([])
s0.append(['Control', 'Da', 'Tiene que dar', 'OK'])
for c in s0[4]: c.font = B; c.fill = AZUL
for nombre, da, debe, ok in controles:
    s0.append([nombre, da, debe, '✓' if ok else '✗ NO CIERRA'])
    s0.cell(s0.max_row, 4).fill = VERDE if ok else ROJO
s0.append([]); s0.append(['Totales', 'Débitos', 'Créditos'])
s0.append(['Movimientos a cargar', tot_deb, tot_cre])
s0.append([]); s0.append(['Deshacer: la app no tiene ninguna fila antes del 01/02/2026, así que se borran todas las de estado «historico» con fecha anterior a esa — y nada más.'])
s0.column_dimensions['A'].width = 95
s1 = wb.create_sheet('01 Filas a cargar')
cols = [k for k in registros[0] if k != '_fila']
s1.append(['Fila de tu planilla'] + cols + ['Control de la cadena (fórmula)'])
for c in s1[1]: c.font = B; c.fill = AZUL
ic = {k: j + 2 for j, k in enumerate(cols)}
from openpyxl.utils import get_column_letter as L
for n, r in enumerate(registros, start=2):
    s1.append([r['_fila']] + [r[k] for k in cols])
    cs, cd, cc = L(ic['saldo']), L(ic['debitos']), L(ic['creditos'])
    prev = f'{SALDO_INICIO}' if n == 2 else f'{cs}{n - 1}'
    s1.cell(n, len(cols) + 2, f'=ROUND({prev}+{cc}{n}-{cd}{n}-{cs}{n},2)')
fin = s1.max_row
s1.append(['TOTAL'] + [f'=SUM({L(ic[k])}2:{L(ic[k])}{fin})' if k in ('debitos', 'creditos') else None for k in cols] +
          [f'=COUNTIF({L(len(cols) + 2)}2:{L(len(cols) + 2)}{fin},"<>0")'])
s1.freeze_panes = 'B2'
try:
    wb.save(SALIDA)
except PermissionError:
    # El usuario lo tiene abierto: se guarda al lado, sin pisar el suyo.
    SALIDA = SALIDA.replace('.xlsx', ' (copia al cargar).xlsx'); wb.save(SALIDA)
print(SALIDA)
for nombre, da, debe, ok in controles: print(('✓' if ok else '✗'), nombre, '→', da)
print('débitos', tot_deb, 'créditos', tot_cre, '| saltos', saltos[:5])

# ── 4. Aplicar (sólo con OK) ────────────────────────────────────────────────────────────────────────────
if APLICAR:
    if not all(c[3] for c in controles): raise SystemExit('✗ Algún control no cierra: no se carga nada.')
    cuerpo = [{k: v for k, v in r.items() if k != '_fila'} for r in registros]
    ids = []
    for d in range(0, len(cuerpo), 200):
        req = urllib.request.Request(env['NEXT_PUBLIC_SUPABASE_URL'] + '/rest/v1/msa_galicia', method='POST',
                                     data=json.dumps(cuerpo[d:d + 200]).encode(), headers={**H_API, 'Prefer': 'return=representation'})
        ids += [x['id'] for x in json.loads(urllib.request.urlopen(req).read())]
    json.dump(ids, open(DIR + 'Paso 1 - ids cargados.json', 'w'))
    print('✓ cargados', len(ids))
