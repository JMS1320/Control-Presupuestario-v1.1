"""🔺 RECIBIDOS · DECLARADOS · EN LA APP — las tres relaciones, MSA ejercicio 25/26. SÓLO LECTURA.

Pedido del usuario 2026-10-25: *«recibidos contra declarados, declarados contra lo que hay en la app y recibidos
contra lo que hay en la app. Que sea claro el reporte en cada cosa»*.

- **Recibidos**: «Mis Comprobantes Recibidos» de ARCA (CSV). Las facturas en USD vienen en dólares con su TC
  aunque el archivo diga «en pesos»: se pasan a pesos.
- **Declarados**: los 12 libros IVA Compras presentados (`libro_iva_periodo_AAAAMM_original.zip`).
- **En la app**: lo que usa el libro diario — jul–dic 2025 del histórico del contador, ene–jun 2026 de las
  facturas cargadas (`comprobantes_arca`), por subdiario. Además se dice si el comprobante existe en la app
  aunque esté en otro subdiario o fuera del ejercicio.

Empareja por datos, de más a menos estricto: punto de venta + número + importe en pesos · número + CUIT ·
CUIT + fecha + importe. Uso:
  python scripts/control_iva_tres_vias.py "<carpeta zips>" "<recibidos.csv>" "<salida.xlsx>"
"""
import sys, os, csv, re
from collections import Counter

_base = os.path.join(os.path.dirname(os.path.abspath(__file__)), "control_iva_presentado_vs_app.py")
_src = open(_base, encoding="utf-8").read().split("# ═══════════════ REPORTE")[0]
carpeta, csv_rec, salida = sys.argv[1], sys.argv[2], sys.argv[3]
sys.argv = [_base, carpeta, salida]
exec(compile(_src, _base, "exec"))        # trae HIST, ARCA (de la base), PRES (los 12 zips), leer_zip, ent, f2
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

NC = {3, 8, 13, 53, 113}                  # notas de crédito: restan
num = lambda s: float(str(s or "0").replace(".", "").replace(",", ".") or 0)
dig = lambda s: re.sub(r"\D", "", str(s or ""))

# ── Recibidos
REC = []
for r in csv.DictReader(open(csv_rec, encoding="utf-8-sig"), delimiter=";"):
    tc = 1 if r["Moneda"] == "$" else (num(r["Tipo Cambio"]) or 1)
    REC.append({"fuente": "R", "fecha": r["Fecha de Emisión"], "tipo": int(r["Tipo de Comprobante"]), "pv": int(r["Punto de Venta"]),
                "nro": int(r["Número Desde"]), "cuit": dig(r["Nro. Doc. Emisor"]), "den": r["Denominación Emisor"],
                "pesos": round(num(r["Imp. Total"]) * tc, 2), "iva": round(num(r["Total IVA"]) * tc, 2),
                "moneda": "USD" if r["Moneda"] != "$" else "$", "usd": num(r["Imp. Total"]) if r["Moneda"] != "$" else None, "tc": tc})
# ── Declarados
DEC = []
for per, (pres, _) in PRES.items():
    for p in pres:
        DEC.append({"fuente": "D", "fecha": p["fecha"], "tipo": p["tipo"], "pv": p["pv"], "nro": p["nro"], "cuit": p["cuit"], "den": p["den"],
                    "pesos": round(p["total"], 2), "iva": round(p["cf"], 2), "per": f"{per[4:]}/{per[:4]}"})
# ── App: lo que usa el libro diario, y todo lo demás (para decir dónde está)
def app_row(a, h):
    if h:
        return {"fuente": "A", "fecha": str(a.get("fecha") or "")[:10], "tipo": ent(str(a.get("tipo") or "").split("-")[0]), "pv": ent(a.get("punto_de_venta")),
                "nro": ent(a.get("numero_desde")), "cuit": dig(a.get("nro_doc_emisor")), "den": a.get("denominacion_emisor"),
                "pesos": round(float(a.get("imp_total") or 0), 2), "sub": (a.get("anio_contable"), a.get("mes_contable")), "tabla": "histórico"}
    usd = (a.get("moneda") or "$") not in ("$", "PES", "ARS")
    return {"fuente": "A", "fecha": str(a.get("fecha_emision") or "")[:10], "tipo": a.get("tipo_comprobante"), "pv": ent(a.get("punto_venta")),
            "nro": ent(a.get("numero_desde")), "cuit": dig(a.get("cuit")), "den": a.get("denominacion_emisor"),
            "pesos": round(float(a.get("imp_total") or 0) * (float(a.get("tipo_cambio") or 1) if usd else 1), 2),
            "sub": (a.get("año_contable"), a.get("mes_contable")), "tabla": "facturas"}
TODA_APP = [app_row(a, True) for a in HIST] + [app_row(a, False) for a in ARCA]
def en_libro_diario(x):
    y, m = x["sub"]
    if not y or not m: return False
    ym = int(y) * 100 + int(m)
    return (x["tabla"] == "histórico" and 202507 <= ym <= 202512) or (x["tabla"] == "facturas" and 202601 <= ym <= 202606)
APP = [x for x in TODA_APP if en_libro_diario(x)]
subtxt = lambda x: f'{x["tabla"]} {int(x["sub"][1]):02d}/{x["sub"][0]}' if x["sub"][0] and x["sub"][1] else f'{x["tabla"]} sin subdiario'

igual = lambda a, b: abs(abs(a) - abs(b)) <= 1
CRIT = [lambda a, b: a["pv"] == b["pv"] and a["nro"] == b["nro"] and igual(a["pesos"], b["pesos"]),
        lambda a, b: a["nro"] == b["nro"] and a["nro"] != 0 and a["cuit"] and a["cuit"] == b["cuit"],
        lambda a, b: a["cuit"] and a["cuit"] == b["cuit"] and a["fecha"] == b["fecha"] and igual(a["pesos"], b["pesos"])]
def emparejar(L1, L2):
    """→ dict id(x de L1) → y de L2 (uno a uno), y los que quedan sueltos de cada lado."""
    libres = list(L2); par = {}
    for c in CRIT:
        for x in L1:
            if id(x) in par: continue
            y = next((y for y in libres if c(x, y)), None)
            if y is not None: par[id(x)] = y; libres.remove(y)
    return par, [x for x in L1 if id(x) not in par], libres
def buscar_en_toda_la_app(x):
    y = next((y for c in CRIT for y in TODA_APP if c(x, y)), None)
    return subtxt(y) if y else "no está en la app"

RD, R_sin_D, D_sin_R = emparejar(REC, DEC)
DA, D_sin_A, A_sin_D = emparejar(DEC, APP)
RA, R_sin_A, A_sin_R = emparejar(REC, APP)
signo = lambda x: -1 if x["tipo"] in NC else 1
plata = lambda L: round(sum(signo(x) * abs(x["pesos"]) for x in L), 2)

# ── Excel
gris = PatternFill("solid", fgColor="D9D9D9"); rojo = PatternFill("solid", fgColor="F8CBAD"); amar = PatternFill("solid", fgColor="FFF2CC"); verde = PatternFill("solid", fgColor="E2EFDA")
B = Font(bold=True); MON = '#,##0.00;-#,##0.00'
wb = Workbook(); r = wb.active; r.title = "Resumen"
def fila(ws, vals, font=None, fill=None, money=()):
    ws.append(vals)
    for c in ws[ws.max_row]:
        if font: c.font = font
        if fill: c.fill = fill
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i in money: ws.cell(ws.max_row, i).number_format = MON
fila(r, ["RECIBIDOS · DECLARADOS · EN LA APP — MSA, ejercicio 25/26 (jul-25 → jun-26)"], Font(bold=True, size=13))
fila(r, ["Recibidos = lo que nos emitieron según ARCA («Mis Comprobantes Recibidos») · Declarados = los 12 libros IVA Compras presentados · En la app = los subdiarios que usa el libro diario (jul–dic 2025 histórico del contador, ene–jun 2026 facturas cargadas). Montos en pesos (lo que vino en USD, × su TC)."])
fila(r, [])
fila(r, ["🔴 LO IMPORTANTE: comprobantes RECIBIDOS que NO se DECLARARON"], Font(bold=True, size=12))
fila(r, ["Fecha", "Emisor", "Tipo", "Nro", "Total", "IVA (crédito fiscal no tomado)", "¿En la app?", "Qué es"], B, gris)
EXPL = {"8768.6": "Es el IVA (21%) de la ND 3-1876 del mismo día ($41.755,19 al 0%, ésa SÍ declarada): Agro Centros facturó el IVA en una ND aparte. En pesos (moneda $, TC 1)."}
for x in sorted(R_sin_D, key=lambda x: x["fecha"]):
    q = EXPL.get(str(round(x["pesos"], 2)), "Factura B: no da crédito fiscal" if x["tipo"] in (6, 7, 8) else ("$0" if x["pesos"] == 0 else ""))
    fila(r, [x["fecha"], x["den"], x["tipo"], f'{x["pv"]}-{x["nro"]}', x["pesos"], x["iva"] if x["tipo"] not in (6, 7, 8) else 0, buscar_en_toda_la_app(x), q], money=(5, 6))
n1 = r.max_row
fila(r, ["TOTAL", "", "", "", plata(R_sin_D), sum(x["iva"] for x in R_sin_D if x["tipo"] not in (6, 7, 8))], B, money=(5, 6))
fila(r, [])
fila(r, ["LAS TRES RELACIONES", "En los dos", "Sólo en el 1º", "$ sólo en el 1º", "Sólo en el 2º", "$ sólo en el 2º", "Control", "Dónde se ve"], B, gris)
def rel(nombre, n1, n2, par, s1, s2, hoja):
    n = r.max_row + 1
    fila(r, [nombre, len(par), len(s1), plata(s1), len(s2), plata(s2), f'=IF(AND(B{n}+C{n}={n1},B{n}+E{n}={n2}),"✓ cierra","❌")', hoja], money=(4, 6))
    for i, v in ((3, len(s1)), (5, len(s2))):
        if v: r.cell(n, i).fill = rojo if i == 3 else amar
rel(f"1 · Recibidos ({len(REC)}) vs Declarados ({len(DEC)})", len(REC), len(DEC), RD, R_sin_D, D_sin_R, "hoja 1")
rel(f"2 · Declarados ({len(DEC)}) vs En la app ({len(APP)})", len(DEC), len(APP), DA, D_sin_A, A_sin_D, "hoja 2")
rel(f"3 · Recibidos ({len(REC)}) vs En la app ({len(APP)})", len(REC), len(APP), RA, R_sin_A, A_sin_R, "hoja 3")
fila(r, [])
fila(r, ["Cómo leer: cada relación tiene su hoja con las dos listas de diferencias. «Todos» tiene cada comprobante con sus tres marcas: un comprobante puede tener hasta dos divergencias (ej.: recibido, no declarado y no cargado)."])
for col, w in zip("ABCDEFGH", [44, 34, 8, 16, 16, 18, 14, 60]): r.column_dimensions[col].width = w

def hoja(nombre, titulo, partes):
    ws = wb.create_sheet(nombre); fila(ws, [titulo], Font(bold=True, size=12))
    for sub, cab, filas, money, fill in partes:
        fila(ws, []); fila(ws, [sub], B)
        fila(ws, cab, B, gris)
        for f in filas: fila(ws, f, money=money, fill=fill)
        if not filas: fila(ws, ["— ninguno —"])
    for i, w in enumerate([11, 40, 7, 16, 15, 16, 26, 50], 1): ws.column_dimensions[chr(64 + i)].width = w
cabR = ["Fecha", "Emisor", "Tipo", "Nro", "Total", "CUIT", "Dónde está en la app", "Nota"]
fr = lambda L, extra: [[x["fecha"], x["den"], x["tipo"], f'{x["pv"]}-{x["nro"]}', x["pesos"], x["cuit"], extra(x), ""] for x in sorted(L, key=lambda x: x["fecha"])]
hoja("1 · Recibidos vs Declarados", "1 · Lo que nos emitieron (ARCA) contra lo que declaramos en los libros IVA", [
    (f"🔴 Recibidos que NO se declararon: {len(R_sin_D)} · ${plata(R_sin_D):,.2f}", cabR, fr(R_sin_D, buscar_en_toda_la_app), (5,), rojo),
    (f"🟡 Declarados que NO figuran como recibidos en ARCA: {len(D_sin_R)} · ${plata(D_sin_R):,.2f} (tickets, servicios públicos, o emitidos a otro CUIT: revisar)",
     ["Fecha", "Emisor", "Tipo", "Nro", "Total", "CUIT", "Período declarado", "Nota"],
     [[x["fecha"], x["den"], x["tipo"], f'{x["pv"]}-{x["nro"]}', x["pesos"], x["cuit"], x["per"], ""] for x in sorted(D_sin_R, key=lambda x: x["fecha"])], (5,), amar)])
hoja("2 · Declarados vs App", "2 · Lo declarado en los libros IVA contra los subdiarios de la app (lo que usa el libro diario)", [
    (f"🔴 Declarados que NO están en el subdiario de la app: {len(D_sin_A)}", ["Fecha", "Emisor", "Tipo", "Nro", "Total", "CUIT", "Período declarado / dónde está en la app", "Nota"],
     [[x["fecha"], x["den"], x["tipo"], f'{x["pv"]}-{x["nro"]}', x["pesos"], x["cuit"], f'{x["per"]} · {buscar_en_toda_la_app(x)}', ""] for x in sorted(D_sin_A, key=lambda x: x["fecha"])], (5,), rojo),
    (f"🟡 En el subdiario de la app pero NO declarados: {len(A_sin_D)} · ${plata(A_sin_D):,.2f}", cabR, fr(A_sin_D, subtxt), (5,), amar),
    ("ℹ️ Declarados en un período y cargados en otro subdiario", ["Fecha", "Emisor", "Tipo", "Nro", "Total", "CUIT", "Declarado → subdiario", "Nota"],
     [[x["fecha"], x["den"], x["tipo"], f'{x["pv"]}-{x["nro"]}', x["pesos"], x["cuit"], f'{x["per"]} → {int(y["sub"][1]):02d}/{y["sub"][0]}', ""]
      for x in DEC if id(x) in DA for y in [DA[id(x)]] if f'{int(y["sub"][1]):02d}/{y["sub"][0]}' != x["per"]], (5,), None)])
hoja("3 · Recibidos vs App", "3 · Lo que nos emitieron (ARCA) contra los subdiarios de la app", [
    (f"🔴 Recibidos que NO están en el subdiario de la app: {len(R_sin_A)} · ${plata(R_sin_A):,.2f}", cabR, fr(R_sin_A, buscar_en_toda_la_app), (5,), rojo),
    (f"🟡 En el subdiario de la app pero NO recibidos según ARCA: {len(A_sin_R)} · ${plata(A_sin_R):,.2f}", cabR, fr(A_sin_R, subtxt), (5,), amar)])

# ── Todos: cada comprobante con sus tres marcas
todos = []
vistos_D, vistos_A = set(), set()
for x in REC:
    d = RD.get(id(x)); a = RA.get(id(x)) or (DA.get(id(d)) if d else None)
    if d: vistos_D.add(id(d))
    if a: vistos_A.add(id(a))
    todos.append((x, "✓", d["per"] if d else "✗", subtxt(a) if a else "✗"))
for d in DEC:
    if id(d) in vistos_D: continue
    a = DA.get(id(d))
    if a: vistos_A.add(id(a))
    todos.append((d, "✗", d["per"], subtxt(a) if a else "✗"))
for a in APP:
    if id(a) not in vistos_A: todos.append((a, "✗", "✗", subtxt(a)))
ws = wb.create_sheet("Todos")
fila(ws, ["Fecha", "Emisor", "Tipo", "Nro", "Total", "Recibido", "Declarado (período)", "En la app (subdiario)", "Divergencias"], B, gris)
for x, rc, dc, ac in sorted(todos, key=lambda t: t[0]["fecha"]):
    div = [k for k, v in (("no recibido", rc), ("no declarado", dc), ("no está en la app", ac)) if v == "✗"]
    fila(ws, [x["fecha"], x["den"], x["tipo"], f'{x["pv"]}-{x["nro"]}', x["pesos"], rc, dc, ac, " · ".join(div)], money=(5,),
         fill=(rojo if len(div) == 2 else amar if div else None))
for i, w in enumerate([11, 40, 7, 16, 15, 9, 14, 22, 36], 1): ws.column_dimensions[chr(64 + i)].width = w
ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
wb.save(salida)
print(f"R {len(REC)} · D {len(DEC)} · A {len(APP)}")
print(f"1 R-D: ambos {len(RD)} · R sin D {len(R_sin_D)} ${plata(R_sin_D):,.2f} · D sin R {len(D_sin_R)} ${plata(D_sin_R):,.2f}")
print(f"2 D-A: ambos {len(DA)} · D sin A {len(D_sin_A)} · A sin D {len(A_sin_D)} ${plata(A_sin_D):,.2f}")
print(f"3 R-A: ambos {len(RA)} · R sin A {len(R_sin_A)} ${plata(R_sin_A):,.2f} · A sin R {len(A_sin_R)} ${plata(A_sin_R):,.2f}")
print("divergencias:", Counter(len([v for v in t[1:] if v == "✗"]) for t in todos))
