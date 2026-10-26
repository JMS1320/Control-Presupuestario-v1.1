"""🧾 ¿FALTAN FACTURAS POR IMPUTAR? — contra lo PRESENTADO a ARCA, no sólo contra la app. SÓLO LECTURA.

Pedido del usuario 2026-10-25: *«faltan facturas imputar? no lo controles sólo contra la app sino contra lo
presentado a ARCA… son 12 comprimidos con los Excels de ARCA de lo presentado en subdiarios»*.

Toma cada comprobante de los 12 libros IVA Compras presentados (jul-25 → jun-26), le busca su par en la app
con el MISMO cruce de `control_iva_presentado_vs_app.py` (jul–dic contra el histórico, ene–jun contra lo cargado
de ARCA, como arma el libro diario) y mira si tiene cuenta contable: **imputada = número + nombre de cuenta**.
Lista también lo que está en la app y no se presentó (para no olvidarlo) y lo presentado que no tiene par.

Uso: python scripts/control_iva_imputacion.py "<carpeta de zips>" "<salida.xlsx>"
"""
import sys, os
from collections import Counter, defaultdict

# Reusa la lectura de los zips y el cruce del control del 21/10 (mismo criterio, un solo lugar).
_base = os.path.join(os.path.dirname(os.path.abspath(__file__)), "control_iva_presentado_vs_app.py")
_src = open(_base, encoding="utf-8").read().split("# ═══════════════ REPORTE")[0]
carpeta, salida = sys.argv[1], sys.argv[2]
sys.argv = [_base, carpeta, salida]
exec(compile(_src, _base, "exec"))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

def cuenta_de(a, fuente):
    if fuente == "hist":
        nombre_c = a.get("cuenta_asignada") or a.get("cuenta_contable")
        return nombre_c, a.get("nro_cuenta")
    return a.get("cuenta_contable"), a.get("nro_cuenta")

filas, resumen, sin_par = [], [], []
for per, (pres, _) in PRES.items():
    anio, mes = int(per[:4]), int(per[4:])
    fuente = "hist" if (anio, mes) < (2026, 1) else "arca"      # diciembre: el libro diario toma el histórico
    base = HIST if fuente == "hist" else ARCA
    ca = "anio_contable" if fuente == "hist" else "año_contable"
    crudos = {a["id"]: a for a in base if a.get(ca) == anio and a.get("mes_contable") == mes}
    app = [norm_app(a, fuente) for a in crudos.values()]
    pares, faltan, sobran = cruzar(pres, app)
    c = Counter()
    for _, p, a in pares:
        nom, nro = cuenta_de(crudos[a["id"]], fuente)
        estado = "✓ imputada" if (nom and nro) else ("⚠ nombre sin número" if nom else ("⚠ número sin nombre" if nro else "❌ SIN CUENTA"))
        c[estado] += 1
        filas.append([f"{mes:02d}/{anio}", "histórico" if fuente == "hist" else "app", estado, p["fecha"], p["den"], p["cuit"], p["pv"], p["nro"], p["total"], nom, nro,
                      crudos[a["id"]].get("estado") if fuente == "arca" else ""])
    for p in faltan:
        c["sin par en la app"] += 1
        sin_par.append([f"{mes:02d}/{anio}", "presentado, sin par en la app", p["fecha"], p["den"], p["cuit"], p["pv"], p["nro"], p["total"]])
    for a in sobran:
        nom, nro = cuenta_de(crudos[a["id"]], fuente)
        sin_par.append([f"{mes:02d}/{anio}", "en la app, NO presentado en ese período", a["fecha"], a["den"], a["cuit"], a["pv"], a["nro"], a["total"],
                        "con cuenta" if (nom and nro) else "SIN CUENTA"])
    resumen.append([f"{mes:02d}/{anio}", "histórico" if fuente == "hist" else "app", len(pres), c["✓ imputada"], c["❌ SIN CUENTA"],
                    c["⚠ nombre sin número"] + c["⚠ número sin nombre"], c["sin par en la app"],
                    round(sum(f[8] for f in filas if f[0] == f"{mes:02d}/{anio}" and not f[2].startswith("✓")), 2)])

gris = PatternFill("solid", fgColor="D9D9D9"); rojo = PatternFill("solid", fgColor="F8CBAD"); amar = PatternFill("solid", fgColor="FFF2CC")
wb = Workbook(); r = wb.active; r.title = "Resumen"
r.append(["¿FALTAN FACTURAS POR IMPUTAR? — Libro IVA Compras PRESENTADO a ARCA vs cuentas en la app · MSA 25/26"]); r["A1"].font = Font(bold=True, size=12)
r.append(["Universo: lo presentado en los 12 libros IVA. Imputada = la factura tiene número y nombre de cuenta. jul–dic se mira en el histórico del contador (así arma el libro diario), ene–jun en la app."])
r.append([])
r.append(["Período", "Dónde se mira", "Presentadas", "Imputadas", "SIN CUENTA", "Cuenta a medias", "Sin par en la app", "$ sin imputar", "Control"])
for c in r[r.max_row]: c.font = Font(bold=True); c.fill = gris
ini = r.max_row + 1
for x in resumen:
    n = r.max_row + 1
    r.append(x + [f'=IF(C{n}=D{n}+E{n}+F{n}+G{n},"✓ cierra","❌")'])
    if x[4] or x[5]: r.cell(n, 5).fill = rojo
fin = r.max_row
r.append(["TOTAL", "", f"=SUM(C{ini}:C{fin})", f"=SUM(D{ini}:D{fin})", f"=SUM(E{ini}:E{fin})", f"=SUM(F{ini}:F{fin})", f"=SUM(G{ini}:G{fin})", f"=SUM(H{ini}:H{fin})"])
for c in r[r.max_row]: c.font = Font(bold=True)
for row in r.iter_rows(min_row=ini): row[7].number_format = '#,##0.00'
for col, w in zip("ABCDEFGHI", [9, 13, 12, 11, 12, 15, 16, 16, 10]): r.column_dimensions[col].width = w

def hoja(nombre, cab, datos, money, anchos):
    ws = wb.create_sheet(nombre); ws.append(cab)
    for c in ws[1]: c.font = Font(bold=True); c.fill = gris
    for d in datos: ws.append(d)
    for row in ws.iter_rows(min_row=2):
        for i in money: row[i].number_format = '#,##0.00'
    for i, w in enumerate(anchos): ws.column_dimensions[chr(65 + i)].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    return ws
hoja("Sin cuenta", ["Período", "Dónde", "Estado", "Fecha", "Emisor", "CUIT", "PV", "Nº", "Total", "Cuenta", "Nº cuenta", "Estado en la app"],
     [f for f in filas if not f[2].startswith("✓")], (8,), [9, 10, 18, 11, 38, 13, 6, 10, 14, 26, 10, 12])
hoja("Sin par", ["Período", "Qué", "Fecha", "Emisor", "CUIT", "PV", "Nº", "Total", "Cuenta"], sin_par, (7,), [9, 34, 11, 38, 13, 6, 10, 14, 11])
hoja("Todas las presentadas", ["Período", "Dónde", "Estado", "Fecha", "Emisor", "CUIT", "PV", "Nº", "Total", "Cuenta", "Nº cuenta", "Estado en la app"],
     filas, (8,), [9, 10, 18, 11, 38, 13, 6, 10, 14, 26, 10, 12])
wb.save(salida)
tot = defaultdict(int)
for x in resumen:
    for k, v in zip(["pres", "imp", "sin", "medias", "sinpar"], x[2:7]): tot[k] += v
print(dict(tot)); [print(x) for x in resumen]
print("sin par (app no presentada):", sum(1 for s in sin_par if s[1].startswith("en la app")), "· sin cuenta de ésas:", sum(1 for s in sin_par if len(s) > 8 and s[8] == "SIN CUENTA"))
