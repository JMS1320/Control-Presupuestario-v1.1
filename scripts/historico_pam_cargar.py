"""
🏦 HISTÓRICO PAM (A-DEC-1010, 2026-10-18): el extracto de PAM —caja de ahorro desde el 02/01/2025 y cuenta corriente
desde su apertura (24/09/2025)— hasta enero 2026, con el MISMO procedimiento que MSA.

    python scripts/historico_pam_cargar.py            → SIMULA (sólo lee) y arma el Excel para auditar
    python scripts/historico_pam_cargar.py --aplicar  → escribe (SÓLO con OK del usuario), foto y control desde la base

PASO 1 · los 361 movimientos entran como «historico», sin categ ni códigos (igual que el Paso 1 de MSA):
  · CA con EL MISMO parser y reglas que el importador (scripts/historico_pam_parseo_ca.mts); CC con sus columnas;
  · `orden` negativo (−N … −1), la cadena de saldos con su control, y que empalme con lo que la app ya tiene;
  · su Detalle/Observaciones en `detalle`; su CATEG / Contable / Interno en `nota_operador`.
PASO 2 · el lote: lo que el usuario marcó en «… TRABAJO v3.xlsx» (11 CATEG). Cada movimiento → una cuota «anterior»
  propia en su template (o el pago de Sueldos), y el movimiento pasa a «conciliado». SIN contable ni interno.
  Reglas del usuario (18/10): templates de PAM siempre (otra empresa sólo con su OK — la fila 274 «Red Vial Lima» es de
  MA y él la aprobó); los nuevos con responsable PAM; los fijos a la campaña de su fecha (2025 / 2026), los abiertos sin
  campaña; su detalle ENTERO en la cuota; los ADICIONALES llevan «Cuota Adicional 2025»; sueldos con «pagado por PAM».
  Lo marcado «no» (PAM ya presentó 2025 y no tiene campañas siguientes) y Expensas Posadas (próxima tanda) quedan históricos.
"""
import json, os, re, sys, datetime, collections, subprocess, calendar, urllib.request, urllib.parse
import openpyxl
from openpyxl.styles import PatternFill, Font

APLICAR = '--aplicar' in sys.argv
RAIZ = 'D:/Users/josem/Documents/Jose/Automatizarr/Claude/Control-Presupuestario-v1.1/'
DIR = RAIZ + '- Comunicacion JMS Claude - Archivos/Balance/Excels trabajando con claude/'
FUENTE = DIR + 'Extracto PAM - Control Presupuestario.xlsx'
DECISIONES = DIR + 'Extracto PAM - Control Presupuestario - TRABAJO v3.xlsx'
# ⚠️ NUNCA se pisa un archivo existente: el 18/10 una corrida pisó la v2 que el usuario ya había marcado y se perdieron
# sus notas. Se toma la primera versión libre (vN), o la que se pida con --salida=… (y si existe, se frena).
_pedida = next((a.split('=', 1)[1] for a in sys.argv if a.startswith('--salida=')), None)
if _pedida:
    SALIDA = _pedida
else:
    import os
    _n = 2
    while os.path.exists(DIR + f'PAM - carga del historico - para revisar v{_n}.xlsx'): _n += 1
    SALIDA = DIR + f'PAM - carga del historico - para revisar v{_n}.xlsx'
FOTO = DIR + 'PAM historico - foto y ids creados.json'
import tempfile
TMP = os.path.join(tempfile.gettempdir(), 'pam_historico_parseo_ca.json')

env = {}
for l in open(RAIZ + '.env.local', encoding='utf-8'):
    if '=' in l and not l.strip().startswith('#'):
        k, v = l.split('=', 1); env[k.strip()] = v.strip().strip('"').strip("'")
KEY = env['SUPABASE_SERVICE_ROLE_KEY']
def llamar(metodo, tabla, params=None, cuerpo=None, schema=None):
    url = env['NEXT_PUBLIC_SUPABASE_URL'] + '/rest/v1/' + tabla + ('?' + urllib.parse.urlencode(params, safe='(),.*:') if params else '')
    h = {'apikey': KEY, 'Authorization': 'Bearer ' + KEY, 'Content-Type': 'application/json', 'Prefer': 'return=representation'}
    if schema: h['Accept-Profile'] = schema; h['Content-Profile'] = schema
    data = json.dumps(cuerpo).encode() if cuerpo is not None else None
    return json.loads(urllib.request.urlopen(urllib.request.Request(url, data=data, method=metodo, headers=h)).read() or '[]')
leer = lambda tabla, params, schema=None: llamar('GET', tabla, params, schema=schema)
r2 = lambda x: round(float(x or 0) + 1e-9, 2)
s = lambda x: '' if x is None else str(x)
def texto(v):
    if v is None or v == '': return None
    if isinstance(v, float) and v.is_integer(): return str(int(v))
    return str(v).strip() or None
money = lambda x: f'{x:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')

# ══ PASO 1 · el extracto ═════════════════════════════════════════════════════════════════════════════════
ws = openpyxl.load_workbook(FUENTE, data_only=True).worksheets[0]
H_CA, H_CC = [i for i in range(1, ws.max_row + 1) if s(ws.cell(i, 1).value).strip() == 'Fecha']
hca = {s(ws.cell(H_CA, c).value).strip(): c for c in range(1, ws.max_column + 1) if ws.cell(H_CA, c).value}
hcc = {s(ws.cell(H_CC, c).value).strip(): c for c in range(1, ws.max_column + 1) if ws.cell(H_CC, c).value}
fil_ca = [i for i in range(H_CA + 1, H_CC) if isinstance(ws.cell(i, 1).value, datetime.datetime)]
fil_cc = [i for i in range(H_CC + 1, ws.max_row + 1) if isinstance(ws.cell(i, 1).value, datetime.datetime)]
fil_ca.sort(reverse=True); fil_cc.sort(reverse=True)          # su planilla va del más nuevo al más viejo

def nota(i, h, cat_col):
    partes = [f"CATEG {texto(ws.cell(i, h[cat_col]).value)}" if texto(ws.cell(i, h[cat_col]).value) else None,
              f"Contable {texto(ws.cell(i, h['Contable']).value)}" if texto(ws.cell(i, h['Contable']).value) else None,
              f"Interno {texto(ws.cell(i, h['Interno']).value)}" if texto(ws.cell(i, h['Interno']).value) else None]
    t = ' · '.join(x for x in partes if x)
    return '📜 Histórico PAM ene25-ene26' + (' — tu planilla: ' + t if t else '')

# CA: el parseo lo hace el parser de la app
json.dump([{'raw': s(ws.cell(i, hca['Movimiento']).value)} for i in fil_ca], open(TMP, 'w', encoding='utf-8'), ensure_ascii=False)
subprocess.run(['npx', 'tsx', 'scripts/historico_pam_parseo_ca.mts', TMP, TMP], cwd=RAIZ, check=True, shell=True,
               stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
parseo = json.load(open(TMP, encoding='utf-8'))

def cadena(filas, h, deb_c, cre_c, armar):
    regs = []
    i0 = filas[0]
    saldo_ini = r2(r2(ws.cell(i0, h['Saldo']).value) - r2(ws.cell(i0, h[cre_c]).value) + r2(ws.cell(i0, h[deb_c]).value))
    prev = saldo_ini
    for k, i in enumerate(filas):
        deb, cre, saldo = r2(ws.cell(i, h[deb_c]).value), r2(ws.cell(i, h[cre_c]).value), r2(ws.cell(i, h['Saldo']).value)
        control = round(prev + cre - deb - saldo, 2)
        r = armar(k, i); r.update({'fecha': ws.cell(i, 1).value.strftime('%Y-%m-%d'), 'debitos': deb, 'creditos': cre, 'saldo': saldo,
                                   'control': control, 'orden': -(len(filas) - k), 'estado': 'historico', 'revisado': False,
                                   'categ': None, 'contable': None, 'interno': None, 'centro_de_costo': None, '_fila': i})
        regs.append(r); prev = saldo
    return saldo_ini, regs

def armar_ca(k, i):
    p = parseo[k]; raw = s(ws.cell(i, hca['Movimiento']).value)
    return {'descripcion': p.get('descripcion') or raw[:200], 'concepto': raw, 'origen': 'CA_GALICIA', 'cuenta': 'PAM Galicia CA',
            'grupo_de_conceptos': p.get('grupo_de_conceptos', ''), 'tipo_de_movimiento': p.get('tipo_de_movimiento', ''),
            **{f'leyendas_adicionales_{n}': p.get(f'leyendas_adicionales_{n}', '') for n in (1, 2, 3, 4)},
            'numero_de_terminal': p.get('numero_de_terminal', ''), 'numero_de_comprobante': p.get('numero_de_comprobante', ''),
            'observaciones_cliente': '', 'detalle': texto(ws.cell(i, hca['Observaciones']).value), 'nota_operador': nota(i, hca, 'Categ')}
def armar_cc(k, i):
    g = lambda n: texto(ws.cell(i, hcc[n]).value)
    return {'descripcion': g('Descripción'), 'origen': g('Origen') or '', 'cuenta': 'pam_galicia_cc',
            'grupo_de_conceptos': g('Grupo de Conceptos'), 'concepto': g('Concepto'), 'numero_de_terminal': g('Número de Terminal'),
            'observaciones_cliente': g('Observaciones Cliente'), 'numero_de_comprobante': g('Número de Comprobante'),
            **{f'leyendas_adicionales_{n}': g(f'Leyendas Adicionales{n}') for n in (1, 2, 3, 4)},
            'tipo_de_movimiento': g('Tipo de Movimiento'), 'detalle': g('Detalle'), 'nota_operador': nota(i, hcc, 'CATEG')}

ini_ca, reg_ca = cadena(fil_ca, hca, 'Débito', 'Crédito', armar_ca)
ini_cc, reg_cc = cadena(fil_cc, hcc, 'Débitos', 'Créditos', armar_cc)
TABLA = {'CA': 'pam_galicia', 'CC': 'pam_galicia_cc'}
controles = []
for cta, ini, regs in (('CA', ini_ca, reg_ca), ('CC', ini_cc, reg_cc)):
    t = TABLA[cta]
    app = leer(t, {'select': 'fecha,saldo,debitos,creditos', 'order': 'orden.asc', 'limit': 1})[0]
    antes = leer(t, {'select': 'id', 'fecha': 'lt.2026-02-01'})
    td, tc = r2(sum(r['debitos'] for r in regs)), r2(sum(r['creditos'] for r in regs))
    empalme = round(regs[-1]['saldo'] + float(app['creditos'] or 0) - float(app['debitos'] or 0) - float(app['saldo']), 2)
    saltos = [r for r in regs if abs(r['control']) > 0.005]
    controles += [
        (f'{cta} · movimientos a cargar ({regs[0]["fecha"]} → {regs[-1]["fecha"]})', len(regs), len(regs), True),
        (f'{cta} · saltos en la cadena de saldos', len(saltos), 0, not saltos),
        (f'{cta} · saldo inicial − débitos + créditos = último saldo', r2(ini - td + tc), regs[-1]['saldo'], abs(ini - td + tc - regs[-1]['saldo']) < 0.01),
        (f'{cta} · empalme con el primer movimiento de la app ({app["fecha"]}, saldo {money(float(app["saldo"]))})', empalme, 0, abs(empalme) < 0.01),
        (f'{cta} · filas que la app YA tiene antes del 01/02/2026', len(antes), 0, not antes)]

# ══ PASO 2 · el lote, desde sus decisiones en la v3 ════════════════════════════════════════════════════════
wv = openpyxl.load_workbook(DECISIONES, data_only=True).worksheets[0]
def cols(hr): return {s(wv.cell(hr, c).value).strip(): c for c in range(1, wv.max_column + 1) if wv.cell(hr, c).value}
vca, vcc = cols(H_CA), cols(H_CC)
reg_por_fila = {r['_fila']: (cta, r) for cta, regs in (('CA', reg_ca), ('CC', reg_cc)) for r in regs}
ELEGIDAS = {'IMP', 'FIJOS BS AS', 'FCI', 'BANC', 'DIST MA', 'DIST HNOS', 'DOM', 'RET', 'AP', 'SUELD', 'JORN'}
decis = {}
for i, (cta, r) in reg_por_fila.items():
    v = vcc if cta == 'CC' else vca
    cat = s(wv.cell(i, v['CATEG' if cta == 'CC' else 'Categ']).value).strip()
    if cat not in ELEGIDAS: continue
    # misma fila, mismo movimiento: la v3 es su planilla + columnas; se verifica por fecha e importe
    assert wv.cell(i, 1).value.strftime('%Y-%m-%d') == r['fecha'] and r2(wv.cell(i, v['Débitos' if cta == 'CC' else 'Débito']).value) == r['debitos'], f'fila {i} no coincide con la v3'
    decis[i] = dict(cat=cat, dest=s(wv.cell(i, v['Propuesta: destino']).value), dec=s(wv.cell(i, v['JMS: decisión']).value).strip(),
                    app_id=s(wv.cell(i, v['App: id']).value))

TODOS = leer('egresos_sin_factura', {'select': '*'})
def tpl(nombre, anio=None, resp='PAM'):
    m = [t for t in TODOS if t['nombre_referencia'] == nombre and t['responsable'] == resp and (anio is None or t['año'] == anio)]
    if len(m) > 1: raise SystemExit(f'⛔ «{nombre}» {anio} de {resp} existe {len(m)} veces')
    return m[0] if m else None

# Templates que se CREAN (o se reusan si ya existen) — todos de PAM
CREAR = {}   # (nombre, año) → (base, cambios)
SIN_CODIGOS = {'codigo_contable': None, 'codigo_interno': None}
def a_crear(nombre, anio, base, cambios):
    if not tpl(nombre, anio): CREAR[(nombre, anio)] = (base, cambios)
for n in ('Comision Caja de Seguridad', 'Metrogas Libertad', 'ABL Libertad Cuota', 'Inmobiliario Complementario Cuota PAM',
          'Inmobiliario Cuota Casco', 'Inmobiliario Cuota Ombu', 'Fondo Educativo'):
    a_crear(n, '2025', tpl(n, '2026'), {'año': '2025', 'activo': False, 'aplica_generacion': False, 'template_origen_id': None})
a_crear('Aporte a MSA', '2026', tpl('Aporte PAM', resp='MSA'), {'responsable': 'PAM', 'categ': 'Aporte a MSA', 'cuit_quien_cobra': '30617786016', 'nombre_quien_cobra': 'MSA'})
a_crear('Retiro de MSA', '2026', tpl('Retiro PAM', resp='MSA'), {'responsable': 'PAM', 'categ': 'Retiro de MSA', 'cuit_quien_cobra': '30617786016', 'nombre_quien_cobra': 'MSA'})
ACA = {'proveedor': 'Automóvil Club Argentino', 'nombre_quien_cobra': 'Automóvil Club Argentino', 'cuit_quien_cobra': None,
       'fecha_primera_cuota': None, 'monto_por_cuota': None, 'observaciones_template': 'Creado en el histórico PAM (18/10): cuota del ACA, débito automático', **SIN_CODIGOS}
a_crear('ACA', '2025', tpl('Metrogas Libertad', '2026'), {**ACA, 'año': '2025', 'activo': False, 'aplica_generacion': False, 'template_origen_id': None})
a_crear('ACA', '2026', tpl('Metrogas Libertad', '2026'), {**ACA, 'template_origen_id': None})
a_crear('Sueldo Domesticas PAM', '2025', tpl('Metrogas Libertad', '2026'),
        {'categ': 'Sueldo Domesticas', 'proveedor': None, 'nombre_quien_cobra': None, 'cuit_quien_cobra': None, 'centro_costo': 'Nazarenas',
         'fecha_primera_cuota': None, 'monto_por_cuota': None, 'año': '2025', 'activo': False, 'aplica_generacion': False,
         'template_origen_id': None, 'observaciones_template': 'Creado en el histórico PAM (18/10): sueldo de la doméstica de Nazarenas (a quién se le pagó, en el detalle)', **SIN_CODIGOS})
RENOMBRAR = ('Retiro MA semestral PAM', 'Retiro MA mensual PAM')

FIJOS = {'Comision Caja de Seguridad', 'Metrogas Libertad', 'AYSA Libertad', 'Expensas Libertad', 'ABL Libertad Cuota',
         'ABL Cochera Posadas Cuota', 'Inmobiliario Complementario Cuota PAM', 'ACA'}
RED_VIAL = [('Red Vial Cuota Casco', 299117.17), ('Red Vial Cuota Entre Rios', 46627.09), ('Red Vial Cuota Tapera 1', 31671.23),
            ('Red Vial Cuota Tapera 2', 31671.23), ('Red Vial Cuota Tapera 3', 33430.75), ('Red Vial Cuota Ombu', 136362.25)]   # su hoja «detalles JMS»
INMOB_ADIC = [('Inmobiliario Cuota Casco', 838905.10), ('Inmobiliario Cuota Ombu', 241456.00), ('Fondo Educativo', 16205.30)]
ADICIONAL = ' — Cuota Adicional 2025'

plan, sueldos, vincular, quedan = [], [], [], collections.Counter()
def cuota(i, nombre, anio, monto=None, extra=''):
    cta, r = reg_por_fila[i]
    plan.append(dict(fila=i, cta=cta, mov=r, nombre=nombre, anio=anio, monto=r2(monto if monto is not None else (r['creditos'] or r['debitos'])),
                     tipo='ingreso' if r['creditos'] else 'egreso', desc=f"{nombre} — {r['detalle'] or r['descripcion']}{extra} — histórico PAM"))
for i, d in sorted(decis.items()):
    cta, r = reg_por_fila[i]; anio = r['fecha'][:4]; dec = d['dec'].lower(); dest = d['dest']
    if dec.startswith('no'): quedan['«no» (PAM ya presentó 2025)'] += 1; continue
    if dec.startswith('proxima'): quedan['próxima tanda (Expensas Posadas)'] += 1; continue
    if dest.startswith('VINCULAR'): vincular.append(dict(fila=i, cta=cta, mov=r, cuota_id=d['app_id'])); continue
    if dest.startswith('SUELDOS') or d['cat'] == 'SUELD': sueldos.append(dict(fila=i, cta=cta, mov=r)); continue
    if i == 335:
        for n, m in RED_VIAL: cuota(i, n, '2025', m)
        continue
    if i == 323:
        for n, m in INMOB_ADIC: cuota(i, n, '2025', m, ADICIONAL)
        continue
    if 'cuota del ACA' in dest: cuota(i, 'ACA', anio); continue
    if 'ABL Posadas deuda' in dest: cuota(i, 'ABL Cochera Posadas Cuota', '2025', extra=' — deuda 2024 y 2025 en una cuota'); continue
    nombre = dest.split('TEMPLATE · ')[1].split(' [')[0].split(' (nuevo')[0]
    nombre = {'Retiro MA semestral PAM': RENOMBRAR[1], 'Sueldo Doméstica NZ': 'Sueldo Domesticas PAM'}.get(nombre, nombre)
    if nombre == 'Expensas Libertad': cuota(i, nombre, '2025')      # v2 aprobada: la expensa es del mes anterior → 2025
    elif nombre in FIJOS: cuota(i, nombre, anio, extra=ADICIONAL if 'adicional' in (r['detalle'] or '').lower() else '')
    elif nombre == 'Sueldo Domesticas PAM':                            # usuario: «anotar los detalles ya que es a quien se le paga»
        a_quien = ' '.join(x for x in (r.get('leyendas_adicionales_1'), r.get('leyendas_adicionales_2')) if x)
        if not a_quien:   # transferencias sin regla de parseo: el destinatario sale del texto del banco (nombre / CUIT)
            ls = [l.strip() for l in (r.get('concepto') or '').splitlines()[1:]]
            a_quien = ' '.join(l.replace('NO  ', 'CUIT ') for l in ls if l and not re.fullmatch(r'\d{22}', l) and 'XXXX' not in l
                               and l not in ('LINK', 'VARIOS') and 'MERCADO LIBRE' not in l)
        cuota(i, nombre, '2025', extra=f' — pagado a: {a_quien}' if a_quien else '')
    else: cuota(i, nombre, None)
assert not [p for p in plan if p['nombre'] == 'Iva Bancario'], 'algo marcado «no» entró al plan'

def resolver(p):
    nombre = RENOMBRAR[0] if p['nombre'] == RENOMBRAR[1] else p['nombre']
    if p['anio'] is None:
        ts = [t for t in TODOS if t['nombre_referencia'] == nombre and t['responsable'] == 'PAM']
        if len(ts) == 1: return ts[0], 'existe'
        if (nombre, '2026') in CREAR: return None, 'se crea (abierto, sin campaña)'
        raise SystemExit(f'⛔ «{nombre}» (abierto) de PAM: {len(ts)} templates')
    t = tpl(nombre, p['anio'])
    if t: return t, 'existe'
    if (nombre, p['anio']) in CREAR: return None, f'se crea ({p["anio"]})'
    raise SystemExit(f'⛔ no existe «{nombre}» {p["anio"]} de PAM y no está en la lista a crear')
for p in plan: p['t'], p['estado_tpl'] = resolver(p)

# Sueldos (rows del 30/12 y el premio del 28/11): JMS y AMS, «pagado por PAM»
emps = {e['nombre']: e for e in leer('empleados', {'select': '*'}, 'sueldos')}
pers = {(p['empleado_id'], p['anio'], p['mes']): p for p in leer('periodos', {'select': '*', 'anio': 'eq.2025'}, 'sueldos')}
camp = [c for c in leer('campanas', {'select': '*'}, 'sueldos') if c['etiqueta'] == '25/26'][0]
for x in sueldos:
    det = (x['mov']['detalle'] or '').upper()
    x['emp'] = 'JMS' if 'JMS' in det else 'AMS'
    x['premio'] = 'PREMIO' in det
    x['mes'] = 11 if x['premio'] else 12
    x['periodo'] = pers.get((emps[x['emp']]['id'], 2025, x['mes']))
assert {(x['emp'], x['mes'], bool(x['periodo'])) for x in sueldos} == {('JMS', 12, False), ('AMS', 12, True), ('JMS', 11, True)}, 'los sueldos no son los esperados'

# Vincular: la cuota existe, está «anterior» y sin movimiento
for v in vincular:
    c = leer('cuotas_egresos_sin_factura', {'select': '*', 'id': 'eq.' + v['cuota_id']})[0]
    t = [t for t in TODOS if t['id'] == c['egreso_id']][0]
    assert abs(float(c['monto']) - (v['mov']['debitos'] or v['mov']['creditos'])) < 0.01 and c['fecha_estimada'] == v['mov']['fecha'], f'cuota {c["id"]} no coincide'
    v['cuota'], v['t'] = c, t

# Control del lote: cada movimiento cierra contra lo que se le asigna
por_mov = collections.defaultdict(float)
for p in plan: por_mov[p['fila']] += p['monto']
no_cierra = [(f, m) for f, m in por_mov.items() if abs(r2(m) - (reg_por_fila[f][1]['debitos'] or reg_por_fila[f][1]['creditos'])) > 0.005]
n_mov = len(por_mov) + len(sueldos) + len(vincular)
controles += [
    ('Lote · movimientos de las 11 CATEG', len(decis), 247, len(decis) == 247),
    ('Lote · se cargan + quedan históricos = total', n_mov + sum(quedan.values()), len(decis), n_mov + sum(quedan.values()) == len(decis)),
    ('Lote · movimientos cuyas cuotas no suman su importe (Red Vial e inmobiliario por partida)', len(no_cierra), 0, not no_cierra),
    ('Lote · templates de OTRA empresa (sólo con OK: fila 274 Red Vial Lima de MA)', sum(1 for v in vincular if v['t']['responsable'] != 'PAM'), 1,
     [v['fila'] for v in vincular if v['t']['responsable'] != 'PAM'] == [274])]

# ══ Excel para auditar ═══════════════════════════════════════════════════════════════════════════════════
AZUL = PatternFill('solid', fgColor='DDEBF7'); VERDE = PatternFill('solid', fgColor='E2EFDA'); ROJO = PatternFill('solid', fgColor='F8CBAD'); B = Font(bold=True)
wb = openpyxl.Workbook(); s0 = wb.active; s0.title = '00 Resumen y controles'
s0.append(['CARGA DEL HISTÓRICO PAM — lo que se escribiría en la app. NO está cargado: es para que lo revises.'])
s0.append([]); s0.append(['Control', 'Da', 'Tiene que dar', 'OK'])
for c in s0[3]: c.font = B; c.fill = AZUL
for n, da, debe, ok in controles:
    s0.append([n, da, debe, '✓' if ok else '✗ NO CIERRA']); s0.cell(s0.max_row, 4).fill = VERDE if ok else ROJO
s0.append([]); s0.append(['Qué se carga', 'Movimientos'])
for c in s0[s0.max_row]: c.font = B; c.fill = AZUL
s0.append(['Paso 1 · extracto como histórico (CA + CC)', len(reg_ca) + len(reg_cc)])
s0.append(['Paso 2 · a template (cuota nueva)', len(por_mov)])
s0.append(['Paso 2 · a una cuota que ya estaba en la app', len(vincular)])
s0.append(['Paso 2 · a Sueldos', len(sueldos)])
for k, v in quedan.items(): s0.append([f'Quedan históricos · {k}', v])
s0.append([]); s0.append(['Templates', 'Qué pasa'])
for c in s0[s0.max_row]: c.font = B; c.fill = AZUL
for (n, a), (b, _) in sorted(CREAR.items()): s0.append([f'{n} [{a}] · PAM', f'se crea (copia de «{b["nombre_referencia"]}» {b["año"]} de {b["responsable"]})'])
s0.append([f'{RENOMBRAR[0]} · PAM', f'se renombra a «{RENOMBRAR[1]}»'])
s0.column_dimensions['A'].width = 100; s0.column_dimensions['B'].width = 60

s3 = wb.create_sheet('03 Templates que se crean')
CAMPOS = [('categ', 'Cuenta (categ)'), ('centro_costo', 'Centro de costo'), ('cuenta_agrupadora', 'Agrupadora (dashboard)'),
          ('responsable_interno', 'Responsable interno'), ('tipo', 'Tipo'), ('tipo_template', 'Fijo / abierto'), ('activo', 'Activo'),
          ('proveedor', 'Proveedor'), ('codigo_contable', 'Código contable'), ('codigo_interno', 'Código interno')]
s3.append(['Template', 'Campaña', 'Empresa', 'Copiado de'] + [n for _, n in CAMPOS])
for c in s3[1]: c.font = B; c.fill = AZUL
for (n, a), (b0, cambios) in sorted(CREAR.items()):
    x = {**b0, **cambios}
    s3.append([n, a, x['responsable'], f"{b0['nombre_referencia']} {b0['año']} ({b0['responsable']})"] + [x.get(k) for k, _ in CAMPOS])
s3.freeze_panes = 'B2'
# ── Por template: arriba el template con sus datos, debajo sus cuotas (pedido del usuario: ver la consistencia) ──
s4 = wb.create_sheet('01 Por template', 1)
GRIS = PatternFill('solid', fgColor='D9D9D9')
CAB = ['Template', 'Campaña', 'Empresa', 'Qué pasa con el template'] + [n for _, n in CAMPOS]
DET = ['', 'Fila de tu planilla', 'Cuenta', 'Fecha', 'Tu detalle', 'Débito', 'Crédito', 'Importe de la cuota', 'Ingreso / egreso', 'Descripción de la cuota']
s4.append(['UNA FILA GRIS POR TEMPLATE (sus datos) y DEBAJO las cuotas que se le cargan. Al final de cada template, su total.'])
grupos = collections.OrderedDict()
for p in sorted(plan, key=lambda p: (p['nombre'], p['anio'] or '', p['mov']['fecha'])):
    grupos.setdefault((p['nombre'], p['anio']), []).append(p)
def datos_tpl(nombre, anio):
    viejo = RENOMBRAR[0] if nombre == RENOMBRAR[1] else nombre
    if (nombre, anio or '2026') in CREAR:
        b0, cambios = CREAR[(nombre, anio or '2026')]
        return {**b0, **cambios, 'año': (cambios.get('año') or b0['año'])}, f"SE CREA (copia de «{b0['nombre_referencia']}» {b0['año']} de {b0['responsable']})"
    t0 = tpl(viejo, anio) if anio else tpl(viejo)
    return t0, ('EXISTE — se RENOMBRA de «' + viejo + '»' if viejo != nombre else 'EXISTE')
def cabecera(nombre, anio, x, que):
    s4.append([]); s4.append(CAB); [setattr(c, 'font', B) for c in s4[s4.max_row]]
    s4.append([nombre, x.get('año') if x else anio, x.get('responsable') if x else '', que] + [x.get(k) if x else None for k, _ in CAMPOS])
    for c in s4[s4.max_row]: c.fill = GRIS; c.font = B
    s4.append(DET); [setattr(c, 'font', B) for c in s4[s4.max_row]]
    return s4.max_row + 1
def total(desde):
    h = s4.max_row
    s4.append(['', 'TOTAL', None, None, None, f'=SUM(F{desde}:F{h})', f'=SUM(G{desde}:G{h})', f'=SUM(H{desde}:H{h})'])
    for c in s4[s4.max_row]: c.font = B
for (nombre, anio), ps in grupos.items():
    x, que = datos_tpl(nombre, anio)
    d0 = cabecera(nombre, anio, x, que)
    for p in ps:
        r = p['mov']
        # en un movimiento por partida (Red Vial / inmobiliario) el débito del banco es el total, no el de esta cuota
        varias = sum(1 for q in plan if q['fila'] == p['fila']) > 1
        s4.append(['', p['fila'], p['cta'], r['fecha'], r['detalle'], None if varias else (r['debitos'] or None), None if varias else (r['creditos'] or None),
                   p['monto'], p['tipo'], p['desc'] + (f"   [el movimiento es de {money(r['debitos'])} en total, por partida]" if varias else '')])
    total(d0)
for v in vincular:
    d0 = cabecera(v['t']['nombre_referencia'], v['t']['año'], v['t'], 'EXISTE — la cuota YA está en la app: sólo se vincula' + ('' if v['t']['responsable'] == 'PAM' else f" ⚠️ template de {v['t']['responsable']}"))
    r = v['mov']
    s4.append(['', v['fila'], v['cta'], r['fecha'], r['detalle'], r['debitos'] or None, r['creditos'] or None, float(v['cuota']['monto']),
               v['cuota'].get('tipo_movimiento'), v['cuota'].get('descripcion') or '(sin descripción; no se cambia)'])
    total(d0)
for e, me in sorted({(x['emp'], x['mes']) for x in sueldos}):
    xs = [x for x in sueldos if (x['emp'], x['mes']) == (e, me)]
    per = xs[0]['periodo']
    s4.append([]); s4.append(['SUELDOS', 'Período', 'Hoy en la app', 'Qué pasa']); [setattr(c, 'font', B) for c in s4[s4.max_row]]
    hoy = (f"monto A {money(float(per['monto_a'] or 0))} · premio {money(float(per['premio'] or 0))}") if per else 'no existe'
    que = 'se CREA con lo pagado' if not per else ('se SUMA al campo Premio' if xs[0]['premio'] else 'se SUMA al monto A')
    s4.append([e, f"{'Nov' if me == 11 else 'Dic'} 2025", hoy, que]); [setattr(c, 'fill', GRIS) for c in s4[s4.max_row]]
    s4.append(DET); [setattr(c, 'font', B) for c in s4[s4.max_row]]
    d0 = s4.max_row + 1
    for x in xs:
        r = x['mov']; s4.append(['', x['fila'], x['cta'], r['fecha'], r['detalle'], r['debitos'], None, r['debitos'], 'pago de sueldo',
                                 ('Premio' if x['premio'] else 'Pago') + f" — {r['detalle']} — pagado por PAM (cta cte)"])
    total(d0)
s4.append([]); s4.append(['', 'TOTAL GENERAL (cuotas + vínculos + sueldos)', None, None, None, None, None, f'=SUMIF(B1:B{s4.max_row},"TOTAL",H1:H{s4.max_row})'])
s4.append(['', 'Control: tiene que ser igual al total de los movimientos del lote', None, None, None, None, None,
           r2(sum((reg_por_fila[f][1]['debitos'] or reg_por_fila[f][1]['creditos']) for f in set(p['fila'] for p in plan) | {v['fila'] for v in vincular} | {x['fila'] for x in sueldos}))])
s4.append(['', 'Diferencia (tiene que dar 0)', None, None, None, None, None, f'=ROUND(H{s4.max_row - 1}-H{s4.max_row},2)'])
for k, w in zip('ABCDEFGHIJKLMN', (34, 14, 10, 30, 34, 14, 14, 16, 16, 30, 16, 16, 14, 14)): s4.column_dimensions[k].width = w
s4.freeze_panes = 'A2'
s2 = wb.create_sheet('01 Lote - una fila por cuota')
s2.append(['Fila de tu planilla', 'Cuenta', 'Fecha', 'Tu detalle', 'Débito', 'Crédito', 'Destino', 'Template', 'Campaña', 'Template: ¿existe?',
           'Importe de la cuota', 'Descripción de la cuota / pago'])
for c in s2[1]: c.font = B; c.fill = AZUL
for p in sorted(plan, key=lambda p: (p['nombre'], p['mov']['fecha'])):
    r = p['mov']; s2.append([p['fila'], p['cta'], r['fecha'], r['detalle'], r['debitos'] or None, r['creditos'] or None, 'Cuota nueva',
                             p['nombre'], p['anio'] or 'abierto', p['estado_tpl'], p['monto'], p['desc']])
for v in vincular:
    r = v['mov']; s2.append([v['fila'], v['cta'], r['fecha'], r['detalle'], r['debitos'] or None, r['creditos'] or None, 'Vincular a cuota existente',
                             v['t']['nombre_referencia'] + ('' if v['t']['responsable'] == 'PAM' else f" ⚠️ de {v['t']['responsable']}"), v['t']['año'], 'existe',
                             float(v['cuota']['monto']), '(la cuota ya existe; no cambia)'])
for x in sueldos:
    r = x['mov']; s2.append([x['fila'], x['cta'], r['fecha'], r['detalle'], r['debitos'], None, 'Sueldos', x['emp'],
                             f"{'nov' if x['mes'] == 11 else 'dic'} 2025", 'período existe' if x['periodo'] else 'período se crea', r['debitos'],
                             ('Premio' if x['premio'] else 'Pago') + f" — {r['detalle']} — pagado por PAM (cta cte)"])
fin = s2.max_row
s2.append(['TOTAL', None, None, None, f'=SUMPRODUCT((A2:A{fin}<>A1:A{fin - 1})*E2:E{fin})', f'=SUMPRODUCT((A2:A{fin}<>A1:A{fin - 1})*F2:F{fin})', None, None, None, None, f'=SUM(K2:K{fin})'])
s2.append(['Control: cuotas − movimientos (tiene que dar 0)', None, None, None, None, None, None, None, None, None, f'=K{fin + 1}-E{fin + 1}-F{fin + 1}'])
s2.freeze_panes = 'A2'
for cta, regs in (('CA', reg_ca), ('CC', reg_cc)):
    sx = wb.create_sheet(f'02 Extracto {cta} a cargar')
    ks = [k for k in regs[0] if k != '_fila']
    sx.append(['Fila de tu planilla'] + ks); [setattr(c, 'font', B) for c in sx[1]]
    for r in regs: sx.append([r['_fila']] + [r[k] for k in ks])
    sx.freeze_panes = 'B2'
import os
if os.path.exists(SALIDA): raise SystemExit(f'⛔ {SALIDA} ya existe: no se pisa (puede tener notas del usuario)')
wb.save(SALIDA)
print(SALIDA)
for n, da, debe, ok in controles: print(('✓' if ok else '✗'), n, '→', da)
print(f'Paso 1: CA {len(reg_ca)} + CC {len(reg_cc)} · Paso 2: {len(plan)} cuotas nuevas para {len(por_mov)} mov · {len(vincular)} a vincular · {len(sueldos)} sueldos · quedan {dict(quedan)}')
print(f'Templates a crear: {len(CREAR)} · a renombrar: 1')
if not APLICAR:
    print('SIMULACIÓN — no se escribió nada. Para aplicar: --aplicar'); sys.exit()
if not all(c[3] for c in controles): raise SystemExit('⛔ algún control no cierra: no se escribe nada.')

# ══ APLICAR ══════════════════════════════════════════════════════════════════════════════════════════════
foto = {'movimientos_creados': {'CA': [], 'CC': []}, 'templates_creados': [], 'renombrado': None, 'cuotas_creadas': [], 'grupos_creados': [],
        'periodos_creados': [], 'periodos_antes': [], 'pagos_creados': []}
guardar = lambda: json.dump(foto, open(FOTO, 'w', encoding='utf-8'), ensure_ascii=False, default=str)
guardar()
# Paso 1
id_de = {}
for cta, regs in (('CA', reg_ca), ('CC', reg_cc)):
    for d in range(0, len(regs), 200):
        lote = regs[d:d + 200]
        ins = llamar('POST', TABLA[cta], None, [{k: v for k, v in r.items() if k != '_fila'} for r in lote])
        for r, x in zip(lote, ins): id_de[r['_fila']] = x['id']; foto['movimientos_creados'][cta].append(x['id'])
        guardar()
# Templates
for (n, a), (b, cambios) in CREAR.items():
    nuevo = {k: v for k, v in b.items() if k not in ('id', 'created_at', 'updated_at')}
    nuevo.update(cambios); nuevo['nombre_referencia'] = n
    t = llamar('POST', 'egresos_sin_factura', None, nuevo)[0]; TODOS.append(t); foto['templates_creados'].append(t['id']); guardar()
t_ren = tpl(RENOMBRAR[0])
llamar('PATCH', 'egresos_sin_factura', {'id': 'eq.' + t_ren['id']}, {'nombre_referencia': RENOMBRAR[1]})
foto['renombrado'] = {'id': t_ren['id'], 'antes': RENOMBRAR[0]}; guardar()
t_ren['nombre_referencia'] = RENOMBRAR[1]
for p in plan:
    if p['t'] is None or p['nombre'] == RENOMBRAR[1]: p['t'] = tpl(p['nombre'], p['anio']) if p['anio'] else tpl(p['nombre'])
# Cuotas (+ grupo de pago cuando un movimiento lleva varias, como hace la app)
grupo = {}
for f, n in collections.Counter(p['fila'] for p in plan).items():
    if n > 1:
        r = reg_por_fila[f][1]; red = f == 335
        g = llamar('POST', 'grupos_pago', None, {'cuit': '30669239765' if red else '30710404611', 'proveedor': 'Municipalidad SP' if red else 'ARBA',
                   'monto_total': r['debitos'], 'estado': 'conciliado', 'observaciones': f"Histórico PAM: {r['detalle']} — una cuota por partida"}, schema='msa')[0]
        grupo[f] = g['id']; foto['grupos_creados'].append(g['id']); guardar()
primera = {}
for p in plan:
    fe = p['mov']['fecha']
    c = llamar('POST', 'cuotas_egresos_sin_factura', None, {'egreso_id': p['t']['id'], 'fecha_estimada': fe, 'fecha_vencimiento': fe, 'monto': p['monto'],
               'estado': 'anterior', 'tipo_movimiento': p['tipo'], 'descripcion': p['desc'], 'grupo_pago_id': grupo.get(p['fila'])})[0]
    foto['cuotas_creadas'].append(c['id']); primera.setdefault(p['fila'], (c['id'], p['t'])); guardar()
def conciliar(f, cambios):
    cta, _ = reg_por_fila[f]
    r = llamar('PATCH', TABLA[cta], {'id': 'eq.' + id_de[f], 'estado': 'eq.historico'}, {'estado': 'conciliado', **cambios})
    if len(r) != 1: raise SystemExit(f'⛔ el movimiento de la fila {f} no se actualizó; lo hecho está en la foto')
for f, (cid, t) in primera.items():
    conciliar(f, {'template_id': t['id'], 'template_cuota_id': grupo.get(f, cid), 'categ': t['categ']})
for v in vincular:
    conciliar(v['fila'], {'template_id': v['t']['id'], 'template_cuota_id': v['cuota']['id'], 'categ': v['t']['categ']})
# Sueldos
ABR = ['Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun', 'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic']
por_per = collections.defaultdict(list)
for x in sueldos: por_per[(x['emp'], x['mes'])].append(x)
for (e, me), xs in sorted(por_per.items()):
    xs.sort(key=lambda x: x['mov']['debitos'])
    tot = r2(sum(x['mov']['debitos'] for x in xs)); per = xs[0]['periodo']
    if per is None:
        per = llamar('POST', 'periodos', None, {'empleado_id': emps[e]['id'], 'campana_id': camp['id'], 'anio': 2025, 'mes': me,
              'fecha_inicio_periodo': f'2025-{me:02d}-01', 'fecha_fin_periodo': f'2025-{me:02d}-{calendar.monthrange(2025, me)[1]:02d}',
              'bruto_calculado': tot, 'sueldo_x_ipc': 0, 'monto_a': tot, 'monto_b': 0, 'anticipos_descontados': tot, 'saldo_pendiente': 0,
              'varios': 0, 'estado': 'historico', 'observaciones': 'Histórico: armado con lo pagado — pagado por PAM (cta cte)'}, schema='sueldos')[0]
        foto['periodos_creados'].append(per['id'])
    else:
        foto['periodos_antes'].append(per); guardar()
        cambio = {'bruto_calculado': r2(float(per['bruto_calculado'] or 0) + tot), 'anticipos_descontados': r2(float(per['anticipos_descontados'] or 0) + tot),
                  'observaciones': (per['observaciones'] or '') + f' · + {money(tot)} pagado por PAM (cta cte)'}
        if xs[0]['premio']: cambio['premio'] = r2(float(per['premio'] or 0) + tot)
        else: cambio['monto_a'] = r2(float(per['monto_a'] or 0) + tot)
        llamar('PATCH', 'periodos', {'id': 'eq.' + per['id']}, cambio, schema='sueldos')
    guardar()
    for k, x in enumerate(xs):
        tipo = 'sueldo' if (per['id'] in foto['periodos_creados'] and k == len(xs) - 1) or x['premio'] else 'anticipo'
        etiqueta = 'Premio' if x['premio'] else ('Pago Saldo' if tipo == 'sueldo' else 'Anticipo')
        pg = llamar('POST', 'pagos', None, {'periodo_id': per['id'], 'empleado_id': emps[e]['id'], 'tipo': tipo, 'fecha': x['mov']['fecha'],
             'monto': x['mov']['debitos'], 'descripcion': f"{etiqueta} {ABR[me - 1]} 2025 — {x['mov']['detalle']} — pagado por PAM (cta cte)",
             'estado': 'anterior', 'medio_pago': 'banco', 'visible_contable': False}, schema='sueldos')[0]
        foto['pagos_creados'].append(pg['id']); guardar()
        conciliar(x['fila'], {'sueldo_pago_id': pg['id'], 'categ': 'Sueldos'})
print(f"✅ aplicado: {sum(len(v) for v in foto['movimientos_creados'].values())} movimientos · {len(foto['templates_creados'])} templates · "
      f"{len(foto['cuotas_creadas'])} cuotas · {len(foto['grupos_creados'])} grupos · {len(foto['pagos_creados'])} pagos de sueldo")

# ══ CONTROL leído de la base ═════════════════════════════════════════════════════════════════════════════
for cta, t in TABLA.items():
    ms = leer(t, {'select': 'id,estado,debitos,creditos,template_cuota_id,sueldo_pago_id,contable,interno,detalle', 'fecha': 'lt.2026-02-01'})
    print(f'CONTROL {cta}: {len(ms)} filas antes del 01/02/2026 · {collections.Counter(m["estado"] for m in ms)} · con códigos: {sum(1 for m in ms if m["contable"] or m["interno"])}')
cs = leer('cuotas_egresos_sin_factura', {'select': 'id,monto,estado', 'id': 'in.(' + ','.join(foto['cuotas_creadas']) + ')'})
print(f"CONTROL cuotas: {len(cs)} creadas, todas «anterior»: {all(c['estado'] == 'anterior' for c in cs)} · suma {money(sum(float(c['monto']) for c in cs))} vs plan {money(sum(p['monto'] for p in plan))}")
