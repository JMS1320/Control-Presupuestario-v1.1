"""
Auditoría de los seguros de Federación Patronal en Excel — A-DAT-87 / A-FEAT-1309.

SÓLO LECTURA sobre la base (GET a la API con la service key). No escribe ningún dato.

Entradas:
  --polizas   JSON que vuelca `JSON=... npm run probar:polizas <carpeta>` (lo leído de cada PDF)
  --pdfs      carpeta con los PDFs copiados (nombre = «<ejercicio>~<mes>~...~<archivo>.pdf»)
  --env       .env.local con NEXT_PUBLIC_SUPABASE_URL y SUPABASE_SERVICE_ROLE_KEY
  --salida    carpeta donde dejar el Excel y las copias de los PDFs (la de comunicación/polizas)

Reglas del usuario para los Excel de trabajo (KNOWLEDGE § «SU planilla completa»):
  gris = su planilla (el extracto, todas las columnas) · azul = lo cargado en la app ·
  verde = lo nuestro (lectura del PDF, propuesta, controles) · débitos y créditos separados ·
  control con fórmulas vivas · si el archivo existe NO se pisa (versión siguiente).
"""
import argparse, json, os, re, shutil, urllib.request, urllib.parse
from datetime import date, datetime, timedelta
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

GRIS = PatternFill("solid", fgColor="E7E6E6")
AZUL = PatternFill("solid", fgColor="DDEBF7")
VERDE = PatternFill("solid", fgColor="E2EFDA")
ROJO = PatternFill("solid", fgColor="F8CBAD")
AMAR = PatternFill("solid", fgColor="FFF2CC")
NEG = Font(bold=True)
FMT = '#,##0.00;[Red]-#,##0.00'


def leer_env(p):
    env = {}
    for l in open(p, encoding="utf-8"):
        if "=" in l and not l.strip().startswith("#"):
            k, v = l.split("=", 1)
            env[k.strip()] = v.strip().strip('"')
    return env


def api(env, tabla, params, esquema="public"):
    url = f"{env['NEXT_PUBLIC_SUPABASE_URL']}/rest/v1/{tabla}?{urllib.parse.urlencode(params)}"
    k = env["SUPABASE_SERVICE_ROLE_KEY"]
    req = urllib.request.Request(url, headers={"apikey": k, "Authorization": "Bearer " + k, "Accept-Profile": esquema})
    with urllib.request.urlopen(req) as r:
        return json.loads(r.read().decode("utf-8"))


def d(s):
    return datetime.strptime(s[:10], "%Y-%m-%d").date() if s else None


def num(x):
    return None if x is None else float(x)


def largo(path):
    """Windows corta en 260 caracteres: las rutas de Drive + scratchpad lo pasan."""
    ab = os.path.abspath(path)
    pre = "\\\\?\\"
    return pre + ab if os.name == "nt" and not ab.startswith(pre) else ab


def libre(path):
    """Primera versión libre: nunca se pisa un archivo que el usuario puede tener abierto o anotado."""
    if not os.path.exists(path):
        return path
    base, ext = os.path.splitext(path)
    n = 2
    while os.path.exists(f"{base}_v{n}{ext}"):
        n += 1
    return f"{base}_v{n}{ext}"


def encabezado(ws, fila, cols):
    for i, (titulo, color) in enumerate(cols, 1):
        c = ws.cell(row=fila, column=i, value=titulo)
        c.fill = color
        c.font = NEG
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(row=fila + 1, column=1)


def ancho(ws, anchos):
    for i, w in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--polizas", required=True)
    ap.add_argument("--pdfs", required=True)
    ap.add_argument("--env", required=True)
    ap.add_argument("--salida", required=True)
    a = ap.parse_args()
    env = leer_env(a.env)
    datos = json.load(open(a.polizas, encoding="utf-8"))
    polizas, cuotas_rep = datos["pdfs"], datos["cuotas"]

    def reparto_de(pol, venc):
        """El reparto de una cuota del consolidado. El plan del endoso y el consolidado pueden diferir
        unos días en el vencimiento (Flota: «7/12 29-01-2026» en el endoso, «01/02/2026» en el
        consolidado): se toma el más cercano dentro de 7 días."""
        por = cuotas_rep.get(pol, {})
        if venc in por:
            return por[venc]
        cerca = [(abs((d(v) - d(venc)).days), v) for v in por]
        cerca = [x for x in cerca if x[0] <= 7]
        return por[min(cerca)[1]] if cerca else {"interno": [], "contable": []}

    # ── 1. Copias de los PDFs (sin pisar) ─────────────────────────────────────────────────────────
    hoy = date.today().isoformat()
    dir_copias = os.path.join(a.salida, f"Copias del subdiario {hoy}")
    os.makedirs(dir_copias, exist_ok=True)
    copia_de = {}
    for p in polizas:
        partes = p["archivo"].split("~")
        nombre = " · ".join(partes[:2] + [partes[-1]]) if len(partes) >= 3 else p["archivo"]
        destino = os.path.join(dir_copias, nombre)
        if not os.path.exists(largo(destino)):
            shutil.copy2(largo(os.path.join(a.pdfs, p["archivo"])), largo(destino))
        copia_de[p["archivo"]] = os.path.join(os.path.basename(dir_copias), nombre)
        p["carpeta"] = "/".join(partes[:-1])
        p["nombre"] = partes[-1]

    # ── 2. La base (sólo lectura) ─────────────────────────────────────────────────────────────────
    filtro_fp = "*patronal*"
    arca = api(env, "comprobantes_arca", {"select": "*", "denominacion_emisor": f"ilike.{filtro_fp}", "order": "fecha_emision"}, "msa")
    hist = api(env, "comprobantes_historico", {"select": "*", "denominacion_emisor": f"ilike.{filtro_fp}", "order": "fecha"}, "msa")
    extracto = api(env, "msa_galicia", {
        "select": "*",
        "or": "(leyendas_adicionales_1.ilike.*patronal*,leyendas_adicionales_2.ilike.*patronal*,leyendas_adicionales_1.ilike.*federac*,leyendas_adicionales_2.ilike.*federac*)",
        "order": "fecha,orden",
    })
    ultimo = api(env, "msa_galicia", {"select": "fecha", "order": "fecha.desc", "limit": "1"})[0]["fecha"]
    primero = api(env, "msa_galicia", {"select": "fecha", "order": "fecha.asc", "limit": "1"})[0]["fecha"]
    cuotas_tpl = api(env, "cuotas_egresos_sin_factura", {"select": "id,fecha_estimada,monto,estado,egreso_id"})
    tpls = {t["id"]: t for t in api(env, "egresos_sin_factura", {"select": "id,nombre_referencia,codigo_contable,codigo_interno,responsable"})}
    cuota_por_id = {c["id"]: c for c in cuotas_tpl}

    def template_de(mov):
        c = cuota_por_id.get(mov.get("template_cuota_id"))
        if c:
            return tpls.get(c["egreso_id"], {}).get("nombre_referencia")
        t = tpls.get(mov.get("template_id"))
        return t.get("nombre_referencia") if t else None

    # ── 3. Comprobante ↔ PDF (por importe y fecha) ────────────────────────────────────────────────
    comprobantes = []
    for h in hist:
        comprobantes.append({"origen": "Histórico (jul–dic 25)", "id": h["id"], "fecha": h["fecha"], "tipo": h.get("tipo"),
            "pv": h.get("punto_de_venta"), "numero": h.get("numero_desde"), "total": num(h.get("imp_total")),
            "neto": num(h.get("imp_neto_gravado")), "iva": num(h.get("iva")), "exento": num(h.get("imp_op_exentas")),
            "perc_iva": num(h.get("percepcion_iva")), "otros": num(h.get("otros_tributos")),
            "cuenta": h.get("cuenta_asignada") or h.get("cuenta_contable"), "nro_cuenta": h.get("nro_cuenta"),
            "estado": None, "pdf_app": None, "empresa": h.get("empresa") or "MSA"})
    for c in arca:
        comprobantes.append({"origen": "Subdiario (ARCA / manual)", "id": c["id"], "fecha": c["fecha_emision"],
            "tipo": f"{c.get('tipo_comprobante')} - {c.get('tipo_comprobante_desc')}", "pv": c.get("punto_venta"),
            "numero": c.get("numero_desde"), "total": num(c.get("imp_total")), "neto": num(c.get("imp_neto_gravado")),
            "iva": num(c.get("iva")), "exento": num(c.get("imp_op_exentas")), "perc_iva": None,
            "otros": num(c.get("otros_tributos")), "cuenta": c.get("cuenta_contable"), "nro_cuenta": c.get("nro_cuenta"),
            "estado": c.get("estado"), "pdf_app": c.get("pdf_drive_url"), "empresa": "MSA"})

    usados = set()
    for cp in comprobantes:
        mejor = None
        for i, p in enumerate(polizas):
            L = p["lectura"]
            if i in usados or L.get("total") is None or cp["total"] is None:
                continue
            if abs(abs(L["total"]) - abs(cp["total"])) > 0.05:
                continue
            dif = abs((d(L.get("fechaEmision")) or d(cp["fecha"])) - d(cp["fecha"])).days
            if dif > 60:
                continue
            if mejor is None or dif < mejor[1]:
                mejor = (i, dif)
        cp["pdf"] = polizas[mejor[0]] if mejor else None
        if mejor:
            usados.add(mejor[0])

    def concepto(L, patron):
        for c in L.get("conceptos", []):
            if re.search(patron, c["concepto"], re.I):
                return c["importe"]
        return None

    # ── 4. Cuotas esperadas de cada póliza (el último consolidado de cada una) ────────────────────
    ultimo_de = {}
    for p in polizas:
        L = p["lectura"]
        if not L.get("consolidado") or not L.get("poliza"):
            continue
        k = L["poliza"]
        clave = (L["consolidado"].get("ultimoEndoso") or 0, L.get("fechaEmision") or "")
        if k not in ultimo_de or clave > ultimo_de[k][0]:
            ultimo_de[k] = (clave, p)
    esperadas = []
    for pol, (_, p) in sorted(ultimo_de.items(), key=lambda x: x[1][1]["lectura"]["consolidado"]["cuotas"][0]["vencimiento"]):
        L = p["lectura"]
        for c in L["consolidado"]["cuotas"]:
            esperadas.append({"poliza": pol, "ramo": L.get("seccion"), "archivo": p["nombre"], "cuota": f"{c['nro']}/{c['de']}",
                              "venc": c["vencimiento"], "importe": c["importe"], "saldo": c.get("saldo"), "p": p})

    movs = [m for m in extracto]
    for m in movs:
        m["_deb"] = num(m.get("debitos")) or 0.0
        m["_cre"] = num(m.get("creditos")) or 0.0
        m["_asig"] = None
    corte = d(ultimo)

    def buscar(monto, venc, credito=False, tol=0.01):
        cands = []
        for m in movs:
            if m["_asig"] is not None:
                continue
            v = m["_cre"] if credito else m["_deb"]
            if abs(v - monto) > tol:
                continue
            dd = (d(m["fecha"]) - venc).days
            if -10 <= dd <= 25:
                cands.append((abs(dd), m))
        return min(cands, key=lambda x: x[0])[1] if cands else None

    for e in esperadas:
        venc = d(e["venc"])
        e["mov"], e["obs"] = None, ""
        if e["importe"] < 0:
            m = buscar(-e["importe"], venc, credito=True)
            if m:
                e["mov"] = m
                e["obs"] = "Reintegro en el banco"
            else:
                e["obs"] = "NC: Federación la COMPENSA con las cuotas siguientes (no hay movimiento bancario)"
            continue
        m = buscar(e["importe"], venc) or (buscar(e["saldo"], venc) if e["saldo"] and e["saldo"] != e["importe"] else None)
        if m:
            e["mov"] = m
            e["obs"] = "OK" if abs(m["_deb"] - e["importe"]) < 0.01 else "Se debitó el SALDO (parte compensada con una NC)"
            continue
        m = buscar(e["importe"], venc, tol=max(5.0, e["importe"] * 0.01))
        if m:
            e["mov"] = m
            e["obs"] = f"⚠️ IMPORTE DISTINTO: la póliza dice {e['importe']:,.2f} y se debitó {m['_deb']:,.2f} — ¿endoso no cargado?"
            continue
        if venc < d(primero):
            e["obs"] = f"Anterior al extracto cargado (empieza el {d(primero).strftime('%d/%m/%Y')})"
        elif venc > corte:
            e["obs"] = f"A vencer (el extracto llega hasta {corte.strftime('%d/%m/%Y')})"
        elif e["saldo"] == 0:
            e["obs"] = "Sin débito y con saldo 0 en el consolidado: cancelada por compensación con una NC"
        else:
            e["obs"] = "🛑 FALTA EL DÉBITO: la cuota venció y no está en el extracto"
    for e in esperadas:
        if e["mov"] is not None:
            e["mov"]["_asig"] = e

    # ── 5. Excel ──────────────────────────────────────────────────────────────────────────────────
    wb = Workbook()
    res = wb.active
    res.title = "Resumen"

    # 5.a Comprobantes
    ws = wb.create_sheet("Comprobantes")
    cols = [("Origen", AZUL), ("Fecha", AZUL), ("Tipo", AZUL), ("Pto. vta.", AZUL), ("Número", AZUL), ("Total", AZUL),
            ("Neto gravado", AZUL), ("IVA", AZUL), ("Exento", AZUL), ("Otros trib.", AZUL), ("Cuenta contable", AZUL),
            ("Estado en la app", AZUL), ("PDF vinculado en la app", AZUL),
            ("PDF: archivo (copia)", VERDE), ("PDF: carpeta en Drive", VERDE), ("PDF: ramo", VERDE), ("PDF: póliza", VERDE),
            ("PDF: endoso", VERDE), ("PDF: total", VERDE), ("PDF: FC / NC", VERDE), ("PDF: neto gravado", VERDE), ("PDF: IVA", VERDE),
            ("PDF: exento", VERDE), ("PDF: cierra consigo mismo", VERDE),
            ("Control: número", VERDE), ("Control: tipo", VERDE), ("Control: total", VERDE), ("Control: neto / IVA", VERDE),
            ("Control: imputación (empresa contable)", VERDE), ("Observación", VERDE)]
    encabezado(ws, 1, cols)
    fila = 2
    errores_comp = 0
    for cp in comprobantes:
        p = cp["pdf"]
        L = p["lectura"] if p else {}
        es_nc = (cp["total"] or 0) < 0 or "Crédito" in str(cp["tipo"]) or str(cp["tipo"]).startswith("3")
        obs = []
        c_num = c_tipo = c_tot = c_neto = c_imp = "—"
        if p:
            nro_ok = str(cp["numero"]) == str(L.get("poliza"))
            fiscal = cp["pv"] and int(cp["pv"]) >= 3 and not nro_ok and len(str(cp["numero"])) <= 6 and str(cp["numero"]) != str(L.get("poliza"))
            c_num = "OK" if nro_ok else ("Número fiscal (no la póliza)" if fiscal else f"⚠️ cargado {cp['numero']}, la póliza es {L.get('poliza')}")
            c_tipo = "OK" if es_nc == bool(L.get("esNotaCredito")) else ("⚠️ es NOTA DE CRÉDITO y está como factura" if L.get("esNotaCredito") else "⚠️ es FACTURA y está como NC")
            c_tot = "OK" if abs(abs(cp["total"]) - abs(L["total"])) <= 0.05 else "⚠️ distinto"
            sg, iv = concepto(L, r"SUBTOTAL GRAVADO"), concepto(L, r"^IVA$")
            if sg is not None and cp["neto"] is not None:
                c_neto = "OK" if abs(abs(cp["neto"]) - abs(sg)) <= 0.05 and abs(abs(cp["iva"] or 0) - abs(iv or 0)) <= 0.05 else f"⚠️ PDF neto {sg:,.2f} / IVA {iv or 0:,.2f}"
            elif sg is None and (cp["neto"] or 0) == 0:
                c_neto = "OK (exento)"
            rc = p.get("repartoContable")
            if rc and rc["partes"]:
                emp = rc["partes"][0]["responsable"] if len(rc["partes"]) == 1 else " / ".join(f"{x['responsable']} {x['proporcion']*100:.0f}%" for x in rc["partes"])
                c_imp = "OK — MSA" if emp == "MSA" else f"⚠️ cargada en MSA y corresponde contable {emp}"
            else:
                c_imp = "OK — MSA"
            # Convención de la carga a mano: punto de venta = endoso, y el endoso 0 va como 1 (ARCA no admite 0).
            if not fiscal and L.get("endoso") is not None and cp["pv"] is not None and int(cp["pv"]) != int(L["endoso"]) and not (int(L["endoso"]) == 0 and int(cp["pv"]) in (0, 1)):
                obs.append(f"ℹ️ punto de venta {cp['pv']} y el endoso es {L['endoso']}")
            if any(str(x).startswith("⚠️") for x in (c_num, c_tipo, c_tot, c_neto, c_imp)):
                errores_comp += 1
        else:
            obs.append("No encontré su PDF en la carpeta del subdiario")
            errores_comp += 1
        valores = [cp["origen"], d(cp["fecha"]), cp["tipo"], cp["pv"], cp["numero"], cp["total"], cp["neto"], cp["iva"], cp["exento"],
                   cp["otros"], cp["cuenta"], cp["estado"], "Sí" if cp["pdf_app"] else ("—" if cp["origen"].startswith("Hist") else "NO"),
                   copia_de.get(p["archivo"]) if p else None, p["carpeta"] if p else None, L.get("seccion"), L.get("poliza"), L.get("endoso"),
                   L.get("total"), ("NC" if L.get("esNotaCredito") else "FC") if p else None, concepto(L, r"SUBTOTAL GRAVADO") if p else None,
                   concepto(L, r"^IVA$") if p else None, concepto(L, r"SUBTOTAL EXENTO") if p else None,
                   ("Sí" if all(x["ok"] for x in p["integridad"]) else "NO") if p else None,
                   c_num, c_tipo, c_tot, c_neto, c_imp, "; ".join(obs)]
        for i, v in enumerate(valores, 1):
            c = ws.cell(row=fila, column=i, value=v)
            if isinstance(v, float):
                c.number_format = FMT
            if isinstance(v, date):
                c.number_format = "dd/mm/yyyy"
            if isinstance(v, str) and (v.startswith("⚠️") or v == "NO"):
                c.fill = ROJO
        fila += 1
    fin_comp = fila - 1
    # PDFs que no casaron con ningún comprobante
    sobrantes = [p for i, p in enumerate(polizas) if i not in usados]
    if sobrantes:
        fila += 1
        ws.cell(row=fila, column=1, value="PDFs del subdiario que no casan con ningún comprobante cargado").font = NEG
        fila += 1
        for p in sobrantes:
            L = p["lectura"]
            ws.cell(row=fila, column=1, value="(sin comprobante)")
            ws.cell(row=fila, column=14, value=copia_de.get(p["archivo"]))
            ws.cell(row=fila, column=15, value=p["carpeta"])
            ws.cell(row=fila, column=16, value=L.get("seccion"))
            ws.cell(row=fila, column=17, value=L.get("poliza"))
            ws.cell(row=fila, column=18, value=L.get("endoso"))
            ws.cell(row=fila, column=19, value=L.get("total")).number_format = FMT
            ws.cell(row=fila, column=30, value="Copia repetida de otro PDF (mismo endoso)" if any(
                q is not p and q["lectura"].get("poliza") == L.get("poliza") and q["lectura"].get("endoso") == L.get("endoso") for q in polizas) else "⚠️ no está cargado")
            fila += 1
    ancho(ws, [20, 11, 18, 8, 11, 14, 14, 13, 13, 11, 20, 11, 10, 42, 34, 20, 11, 7, 14, 6, 14, 13, 13, 9, 30, 30, 10, 26, 34, 34])

    # 5.b Cuotas de las pólizas vs. extracto
    wc = wb.create_sheet("Cuotas vs extracto")
    cols = [("Póliza", VERDE), ("Ramo", VERDE), ("Último PDF leído", VERDE), ("Cuota", VERDE), ("Vencimiento", VERDE),
            ("Importe cuota", VERDE), ("Saldo según Federación", VERDE),
            ("Extracto: fecha", GRIS), ("Extracto: débito", GRIS), ("Extracto: crédito", GRIS), ("Estado en la app", AZUL),
            ("Template conciliado", AZUL), ("Diferencia", VERDE), ("Resultado", VERDE), ("Reparto INTERNO de la cuota", VERDE), ("Reparto CONTABLE de la cuota", VERDE)]
    encabezado(wc, 1, cols)
    fila = 2
    for e in esperadas:
        m = e["mov"]
        valores = [e["poliza"], e["ramo"], e["archivo"], e["cuota"], d(e["venc"]), e["importe"], e["saldo"],
                   d(m["fecha"]) if m else None, m["_deb"] if m else None, m["_cre"] if m else None,
                   m.get("estado") if m else None, template_de(m) if m else None, None, e["obs"],
                   " / ".join(f"{x['responsable']} {x['importe']:,.2f}" for x in reparto_de(e["poliza"], e["venc"])["interno"]),
                   " / ".join(f"{x['responsable']} {x['importe']:,.2f}" for x in reparto_de(e["poliza"], e["venc"])["contable"])]
        for i, v in enumerate(valores, 1):
            c = wc.cell(row=fila, column=i, value=v)
            if isinstance(v, float):
                c.number_format = FMT
            if isinstance(v, date):
                c.number_format = "dd/mm/yyyy"
        if m:
            wc.cell(row=fila, column=13, value=f"=IF(F{fila}<0,-F{fila}-J{fila},F{fila}-I{fila})").number_format = FMT
        if e["obs"].startswith(("🛑", "⚠️")):
            wc.cell(row=fila, column=14).fill = ROJO
        elif e["obs"].startswith("A vencer"):
            wc.cell(row=fila, column=14).fill = AMAR
        fila += 1
    fin_cuotas = fila - 1
    ancho(wc, [11, 22, 36, 7, 11, 14, 14, 11, 14, 14, 11, 26, 12, 70, 44, 30])

    # 5.c Extracto completo de Federación (SU planilla, todas las columnas) + lo nuestro
    wx = wb.create_sheet("Extracto Federacion")
    columnas = [k for k in extracto[0].keys() if not k.startswith("_")] if extracto else []
    cols = [(f"Extracto: {k}", GRIS) for k in columnas] + [("Template (app)", AZUL), ("Propuesta: póliza", VERDE),
            ("Propuesta: cuota", VERDE), ("Propuesta: contable", VERDE), ("Propuesta: interno", VERDE), ("Observación", VERDE)]
    encabezado(wx, 1, cols)
    fila = 2
    sin_poliza = 0
    for m in movs:
        e = m["_asig"]
        obs, pol, cuo, cont, inte = "", None, None, None, None
        if e:
            pol, cuo = e["poliza"], e["cuota"]
            # El reparto de ESTA cuota sumando sus endosos (cada endoso trae el costo de sus bienes).
            rep = reparto_de(e["poliza"], e["venc"])
            fmt = lambda partes: (partes[0]["responsable"] if len(partes) == 1 else " / ".join(f"{x['responsable']} {x['importe']:,.2f}" for x in partes)) if partes else "MSA"
            cont, inte = fmt(rep["contable"]), fmt(rep["interno"])
            procesado = m.get("estado") in ("conciliado", "auditar")
            falta = []
            if cont != "MSA" and cont.find("/") < 0 and (m.get("contable") or "").upper().find(cont) < 0:
                falta.append(f"contable {cont}")
            if inte != "MSA" and "/" in inte and (m.get("interno") or "").lower() != "desglosar":
                falta.append("desglosar el interno")
            elif inte != "MSA" and "/" not in inte and inte not in (m.get("interno") or ""):
                falta.append(f"interno {inte}")
            if not falta:
                obs = "OK"
            elif procesado:
                obs = "⚠️ procesado sin: " + ", ".join(falta)
            else:
                obs = f"Sin procesar ({m.get('estado')}): al conciliarlo, poner " + ", ".join(falta)
        else:
            sin_poliza += 1
            if abs(m["_deb"] - 12902) < 1:
                obs = "🛑 Débito SIN PÓLIZA en el subdiario 25/26 — por importe y fecha sería la AP 9536537 (anterior de la 10413411, Rojas Celia), facturada antes de jul-25"
            elif m["_cre"] > 0:
                obs = "Reintegro de Federación — ¿siniestro o devolución? no hay NC que lo explique"
            elif m["fecha"] < "2025-07-15":
                obs = "Cuota de una póliza del ejercicio anterior (Flota 32436252)"
            else:
                obs = "🛑 Débito sin póliza identificada"
        valores = [m.get(k) for k in columnas] + [template_de(m), pol, cuo, cont, inte, obs]
        for i, v in enumerate(valores, 1):
            k = columnas[i - 1] if i <= len(columnas) else None
            if k in ("debitos", "creditos", "saldo", "control", "orden") and v is not None:
                v = float(v)
            if k == "fecha" and v:
                v = d(v)
            c = wx.cell(row=fila, column=i, value=v)
            if isinstance(v, float):
                c.number_format = FMT
            if isinstance(v, date):
                c.number_format = "dd/mm/yyyy"
            if isinstance(v, str) and v.startswith(("⚠️", "🛑")):
                c.fill = ROJO
        fila += 1
    fin_ext = fila - 1
    col_deb = get_column_letter(columnas.index("debitos") + 1)
    col_cre = get_column_letter(columnas.index("creditos") + 1)
    col_pol = get_column_letter(len(columnas) + 2)
    for i in range(1, len(cols) + 1):
        wx.column_dimensions[get_column_letter(i)].width = 14
    wx.column_dimensions[get_column_letter(len(cols))].width = 70

    # 5.d Bienes asegurados y reparto
    wb_ = wb.create_sheet("Bienes y reparto")
    cols = [("Póliza", VERDE), ("Ramo", VERDE), ("Endoso", VERDE), ("PDF", VERDE), ("Bien", VERDE), ("Patente / DNI", VERDE),
            ("Movimiento (personas)", VERDE), ("Prima del bien", VERDE), ("Responsable contable", VERDE), ("Responsable interno", VERDE),
            ("Por qué", VERDE)]
    encabezado(wb_, 1, cols)
    fila = 2
    for p in polizas:
        L = p["lectura"]
        for r in p.get("responsables", []):
            for i, v in enumerate([L.get("poliza"), L.get("seccion"), L.get("endoso"), p["nombre"], r.get("bien"), r.get("clave"),
                                   r.get("movimiento"), r.get("prima"), r.get("contable"), r.get("interno"), r.get("motivo")], 1):
                c = wb_.cell(row=fila, column=i, value=v)
                if isinstance(v, float):
                    c.number_format = FMT
            fila += 1
    ancho(wb_, [11, 22, 7, 40, 40, 14, 12, 14, 12, 12, 46])

    # 5.e Resumen con CONTROL de fórmulas vivas
    res["A1"] = "Auditoría de seguros — Federación Patronal (MSA)"
    res["A1"].font = Font(bold=True, size=14)
    res["A2"] = f"Generado {datetime.now().strftime('%d/%m/%Y %H:%M')} · extracto hasta {corte.strftime('%d/%m/%Y')} · sólo lectura, no se modificó ningún dato"
    filas = [
        ("COMPROBANTES", None),
        ("Comprobantes cargados (histórico + subdiario)", f"=COUNTA(Comprobantes!A2:A{fin_comp})"),
        ("… con su PDF encontrado en el subdiario", f"=COUNTA(Comprobantes!N2:N{fin_comp})"),
        ("… con alguna observación (⚠️)", errores_comp),
        ("… sin PDF vinculado en la app", f'=COUNTIF(Comprobantes!M2:M{fin_comp},"NO")'),
        ("Total cargado (neto de NC)", f"=SUM(Comprobantes!F2:F{fin_comp})"),
        ("Total según los PDF (neto de NC)", f"=SUM(Comprobantes!S2:S{fin_comp})"),
        ("Diferencia cargado − PDF", "=B9-B10"),
        ("… explicada por NC cargadas como factura (cuentan dos veces su importe)", f'=2*SUMIF(Comprobantes!Z2:Z{fin_comp},"⚠️ es NOTA*",Comprobantes!F2:F{fin_comp})'),
        ("Diferencia SIN explicar (tiene que dar 0)", "=B11-B12"),
        ("", None),
        ("EXTRACTO (movimientos con leyenda FEDERACION PATRO)", None),
        ("Movimientos", f"=COUNTA('Extracto Federacion'!A2:A{fin_ext})"),
        ("Débitos", f"=SUM('Extracto Federacion'!{col_deb}2:{col_deb}{fin_ext})"),
        ("Créditos (reintegros)", f"=SUM('Extracto Federacion'!{col_cre}2:{col_cre}{fin_ext})"),
        ("Débitos asignados a una cuota de póliza", f"=SUMIF('Extracto Federacion'!{col_pol}2:{col_pol}{fin_ext},\"<>\",'Extracto Federacion'!{col_deb}2:{col_deb}{fin_ext})"),
        ("Débitos SIN póliza", f"=SUMIF('Extracto Federacion'!{col_pol}2:{col_pol}{fin_ext},\"\",'Extracto Federacion'!{col_deb}2:{col_deb}{fin_ext})"),
        ("Control: débitos − asignados − sin póliza (tiene que dar 0)", "=B17-B19-B20"),
        ("", None),
        ("CUOTAS DE LAS PÓLIZAS (último pago consolidado de cada una)", None),
        ("Cuotas", f"=COUNTA('Cuotas vs extracto'!A2:A{fin_cuotas})"),
        ("… OK", f"=COUNTIF('Cuotas vs extracto'!N2:N{fin_cuotas},\"OK\")"),
        ("… FALTA EL DÉBITO", f"=COUNTIF('Cuotas vs extracto'!N2:N{fin_cuotas},\"🛑*\")"),
        ("… importe distinto", f"=COUNTIF('Cuotas vs extracto'!N2:N{fin_cuotas},\"⚠️*\")"),
        ("… a vencer", f"=COUNTIF('Cuotas vs extracto'!N2:N{fin_cuotas},\"A vencer*\")"),
        ("… NC compensadas, saldos o anteriores al extracto", "=B24-B25-B26-B27-B28"),
    ]
    for i, (t, f) in enumerate(filas, 4):
        res.cell(row=i, column=1, value=t)
        if f is None and t:
            res.cell(row=i, column=1).font = NEG
        if f is not None:
            c = res.cell(row=i, column=2, value=f)
            c.number_format = FMT if ("Total" in t or "Débitos" in t or "Créditos" in t or "Diferencia" in t or "Control" in t) else "0"
    res.column_dimensions["A"].width = 62
    res.column_dimensions["B"].width = 18
    res["A32"] = "Colores: gris = el extracto tal cual (tu planilla) · azul = lo cargado en la app · verde = lo leído de los PDF y lo propuesto · rojo = a revisar"
    res["A33"] = f"Los PDF copiados están en la carpeta «{os.path.basename(dir_copias)}», al lado de este Excel."

    salida = libre(os.path.join(a.salida, f"Auditoria seguros Federacion Patronal {hoy}.xlsx"))
    wb.save(salida)
    print(salida)
    print(f"comprobantes {len(comprobantes)} · con observación {errores_comp} · cuotas {len(esperadas)} · extracto {len(movs)} · sin póliza {sin_poliza} · pdfs sin comprobante {len(sobrantes)}")


if __name__ == "__main__":
    main()
