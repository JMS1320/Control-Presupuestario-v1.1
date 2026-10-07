"""
📑 Libro de templates sugeridos — VERSIÓN 2, con la planilla del usuario COMPLETA (A-DAT-68, 2026-10-07).

Pedido del usuario: «al quitar los datos que te di me cuesta más auditar… siempre que se usa un movimiento,
que vaya con la info completa: toda mi BBDD tal cual más las columnas tuyas como anexas; mis columnas
agregadas a tus reportes, agregalas; nombre nuevo para no confundir».

    python scripts/sugeridos_v2.py

- A la IZQUIERDA: «Fila planilla» + todas las columnas de su planilla, en su orden y con sus nombres.
- A la DERECHA, en VERDE: las columnas de la propuesta («Propuesta: …»), tomadas del libro v1 tal cual
  (no se recalcula ninguna sugerencia).
- Más a la derecha, en AMARILLO: las columnas que el usuario agregó a mano en el v1 («JMS: …»), copiadas
  por movimiento.
- Subtotales al pie y el CONTROL del resumen con fórmulas (como lo armó él).
El v1 NO se toca.
"""
import openpyxl, collections, datetime, sys
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter as L

DIR = 'D:/Users/josem/Documents/Jose/Automatizarr/Claude/Control-Presupuestario-v1.1/- Comunicacion JMS Claude - Archivos/Balance/'
PLANILLA = DIR + 'extracto desde julio del 2025 hasta enero 26 parcial. galicia MSA cta cte pesos.xlsx'
V1 = DIR + 'Extracto_2025-07_a_2026-01_TEMPLATES_SUGERIDOS.xlsx'
V2 = DIR + 'Extracto_2025-07_a_2026-01_TEMPLATES_SUGERIDOS_v2_completo.xlsx'

VERDE = PatternFill('solid', fgColor='E2EFDA'); AMARILLO = PatternFill('solid', fgColor='FFF2CC'); GRIS = PatternFill('solid', fgColor='EDEDED')
NEGRITA = Font(bold=True)
MONEDA = '#,##0.00'

def r2(x):
    try: return round(float(x), 2)
    except (TypeError, ValueError): return 0.0
def fecha_iso(v):
    if isinstance(v, datetime.datetime): return v.strftime('%Y-%m-%d')
    return str(v or '').strip()[:10]

# ── 1 · La planilla del usuario, completa ──
wp = openpyxl.load_workbook(PLANILLA, data_only=True)
sp = wp.worksheets[0]
hdr_pl = [c.value for c in sp[3]]
while hdr_pl and hdr_pl[-1] in (None, ''): hdr_pl.pop()
ix = {str(h).strip(): i for i, h in enumerate(hdr_pl) if h not in (None, '')}
filas_pl = []
for n, row in enumerate(sp.iter_rows(min_row=4, values_only=True), start=4):
    vals = list(row[:len(hdr_pl)])
    if not isinstance(vals[ix['Fecha']], datetime.datetime): continue
    filas_pl.append((n, vals))
print('planilla:', len(filas_pl), 'movimientos ·', len(hdr_pl), 'columnas')

def clave(fecha, desc, deb, cre):
    return (fecha_iso(fecha), str(desc or '').strip(), r2(deb), r2(cre))

# ── 2 · El v1: mis columnas (G..K) y las suyas (L..) por movimiento ──
w1 = openpyxl.load_workbook(V1, data_only=True)
def leer_v1(hoja, hasta_fila=None):
    ws = w1[hoja]
    hdr = [c.value for c in ws[1]]
    out = collections.defaultdict(list)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0] or not str(row[0])[:4].isdigit(): continue
        out[clave(row[0], row[1], row[2], row[3])].append(list(row))
    return hdr, out
hdr02, v1_02 = leer_v1('02 Movimientos')
hdr03, v1_03 = leer_v1('03 Sin sugerencia')
MIAS = ['Template sugerido', 'Cuenta', 'Confianza', 'Cómo se sugirió', '¿Cargar?']
i_mias = [hdr02.index(h) for h in MIAS]
# Las columnas que agregó el usuario en 03 (de la L en adelante), aunque no tengan título.
ultima03 = max((i for r in v1_03.values() for f in r for i, v in enumerate(f) if v not in (None, '')), default=10)
suyas03 = list(range(11, max(ultima03, len(hdr03) - 1) + 1))
tit_suyas = [f"JMS: {hdr03[i] if i < len(hdr03) and hdr03[i] else '(sin título ' + L(i + 1) + ')'}" for i in suyas03]

# ── 3 · Unir cada movimiento de la planilla con su fila del v1 ──
usados02 = collections.Counter(); usados03 = collections.Counter()
movs = []
sin_match = []
for n, vals in filas_pl:
    k = clave(vals[ix['Fecha']], vals[ix['Descripción']], vals[ix['Débitos']], vals[ix['Créditos']])
    lista02 = v1_02.get(k, [])
    if usados02[k] >= len(lista02): sin_match.append((n, k)); continue
    f02 = lista02[usados02[k]]; usados02[k] += 1
    mias = [f02[i] for i in i_mias]
    suyas = [None] * len(suyas03)
    lista03 = v1_03.get(k, [])
    en03 = False
    if not mias[0] and usados03[k] < len(lista03):
        f03 = lista03[usados03[k]]; usados03[k] += 1; en03 = True
        suyas = [f03[i] if i < len(f03) else None for i in suyas03]
    movs.append({'fila': n, 'vals': vals, 'mias': mias, 'suyas': suyas, 'en03': en03})
print('unidos:', len(movs), '· sin pareja en el v1:', len(sin_match), sin_match[:5])
if sin_match: sys.exit('⚠ hay movimientos de la planilla que no están en el v1 — no se genera nada')

# ── 4 · Armar el v2 ──
wb = openpyxl.Workbook(); wb.remove(wb.active)
def hoja_movs(nombre, lista, con_suyas):
    ws = wb.create_sheet(nombre)
    cab = ['Fila planilla'] + [str(h) if h is not None else '' for h in hdr_pl] + [f'Propuesta: {h}' for h in MIAS] + (tit_suyas if con_suyas else [])
    ws.append(cab)
    n_pl = 1 + len(hdr_pl); n_mias = len(MIAS)
    for j, c in enumerate(ws[1], start=1):
        c.font = NEGRITA
        if n_pl < j <= n_pl + n_mias: c.fill = VERDE
        elif j > n_pl + n_mias: c.fill = AMARILLO
        else: c.fill = GRIS
    for m in lista:
        ws.append([m['fila']] + m['vals'] + m['mias'] + (m['suyas'] if con_suyas else []))
    ult = ws.max_row
    cD, cC, cCargar = L(2 + ix['Débitos']), L(2 + ix['Créditos']), L(n_pl + 5)
    for j in range(1, ws.max_column + 1):
        col = L(j)
        for i in range(2, ult + 1):
            cel = ws[f'{col}{i}']
            if isinstance(cel.value, datetime.datetime): cel.number_format = 'DD/MM/YYYY'
            elif isinstance(cel.value, float): cel.number_format = MONEDA
            if n_pl < j <= n_pl + n_mias: cel.fill = VERDE
            elif j > n_pl + n_mias: cel.fill = AMARILLO
    ws.append([])
    ws.append(['TOTAL'] + [None] * (ix['Débitos']) + [f'=SUM({cD}2:{cD}{ult})', f'=SUM({cC}2:{cC}{ult})'])
    fila_total = ws.max_row
    ws.append(['  de esos, NO se cargan (no son gasto)'] + [None] * (ix['Débitos']) +
              [f'=SUMIFS({cD}2:{cD}{ult},{cCargar}2:{cCargar}{ult},"NO*")', f'=SUMIFS({cC}2:{cC}{ult},{cCargar}2:{cCargar}{ult},"NO*")'])
    fila_no = ws.max_row
    for f in (fila_total, fila_no):
        ws[f'A{f}'].font = NEGRITA
        ws[f'{cD}{f}'].number_format = MONEDA; ws[f'{cC}{f}'].number_format = MONEDA
    ws.freeze_panes = 'B2'
    for j in range(1, ws.max_column + 1): ws.column_dimensions[L(j)].width = 14
    ws.column_dimensions[L(2 + ix['Descripción'])].width = 36
    return ws, fila_total, fila_no, cD, cC

# 00 Resumen: se copia el texto del v1 y se rehace su CONTROL con fórmulas a los nuevos subtotales.
r1 = w1['00 Resumen']
ws00 = wb.create_sheet('00 Resumen')
for row in r1.iter_rows(min_row=1, max_row=12, values_only=True): ws00.append(list(row))
ws00['A13'] = 'v2: cada movimiento con la planilla COMPLETA del usuario (gris), la propuesta (verde) y las columnas agregadas por el usuario (amarillo).'

# 01 Por template y mes: tal cual el v1 (valores), con su columna «Tanda».
r01 = w1['01 Por template y mes']
ws01 = wb.create_sheet('01 Por template y mes')
for row in r01.iter_rows(values_only=True): ws01.append(list(row))
for c in ws01[1]: c.font = NEGRITA
tcol = [c.column for c in ws01[1] if c.value and 'tanda' in str(c.value).lower()]
if tcol: ws01.cell(1, tcol[0]).fill = AMARILLO
ult01 = max(i for i in range(2, ws01.max_row + 1) if ws01.cell(i, 1).value not in (None, ''))
for j in range(3, 11):
    ws01.cell(ult01 + 3, j).value = f'=SUM({L(j)}2:{L(j)}{ult01})'
for row in ws01.iter_rows(min_row=2):
    for c in row:
        if isinstance(c.value, (int, float)): c.number_format = MONEDA

ws02, t02, no02, cD, cC = hoja_movs('02 Movimientos', movs, True)
ws03, t03, no03, _, _ = hoja_movs('03 Sin sugerencia', [m for m in movs if not m['mias'][0]], True)

# El CONTROL del usuario, con fórmulas a los subtotales del v2
fila_total_sug = next(i for i in range(1, 13) if 'Total de gasto sugerido' in str(ws00.cell(i, 1).value or ''))
ws00['A16'] = 'CONTROL'; ws00['B16'] = 'debitos'; ws00['C16'] = 'creditos'
ws00['A17'] = 'movimientos'; ws00['B17'] = f"='02 Movimientos'!{cD}{t02}"; ws00['C17'] = f"='02 Movimientos'!{cC}{t02}"
ws00['A18'] = 'sin sug'; ws00['B18'] = f"='03 Sin sugerencia'!{cD}{t03}"; ws00['C18'] = f"='03 Sin sugerencia'!{cC}{t03}"
ws00['A19'] = 'los que dice que no x no ser gasto'; ws00['B19'] = f"='02 Movimientos'!{cD}{no02}"; ws00['C19'] = f"='02 Movimientos'!{cC}{no02}"
ws00['A20'] = 'sugeridos x dif'; ws00['B20'] = '=B17-B18-B19'; ws00['C20'] = '=C17-C18-C19'
ws00['A21'] = 'control'; ws00['B21'] = f'=B{fila_total_sug}-B20'
for c in ('A16', 'A21'): ws00[c].font = NEGRITA
for f in range(17, 22):
    for col in 'BC': ws00[f'{col}{f}'].number_format = MONEDA
ws00.column_dimensions['A'].width = 48; ws00.column_dimensions['B'].width = 18; ws00.column_dimensions['C'].width = 18

wb.save(V2)

# Control en Python (los mismos números que van a dar las fórmulas)
deb = lambda ms: round(sum(r2(m['vals'][ix['Débitos']]) for m in ms), 2)
cre = lambda ms: round(sum(r2(m['vals'][ix['Créditos']]) for m in ms), 2)
no = [m for m in movs if str(m['mias'][4] or '').startswith('NO')]
s03 = [m for m in movs if not m['mias'][0]]
sug = round(deb(movs) - deb(s03) - deb(no), 2)
print('v2:', V2)
print(f'movimientos {len(movs)} · 03 {len(s03)} (con sus columnas copiadas {sum(1 for m in movs if m["en03"])}) · no cargar {len(no)}')
print(f'sugeridos por diferencia {sug:,.2f} vs resumen {ws00.cell(fila_total_sug, 2).value:,.2f} → control {round(ws00.cell(fila_total_sug, 2).value - sug, 2)}')
print('columnas suyas copiadas:', tit_suyas)
