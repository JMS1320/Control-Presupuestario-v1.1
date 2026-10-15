"""
📒 Excel de TRABAJO del histórico MSA (A-DAT-68, método nuevo 2026-10-15). SÓLO LECTURA de la base.

Toma el Excel del usuario «Extracto MSA - Control Presupuestario.xlsx» (su planilla, ene-25 → ene-26) y
guarda una COPIA con versión («… - TRABAJO vN.xlsx»): su hoja queda igual y se agregan columnas a la derecha.
La v1 sólo marca LO QUE YA ESTÁ EN LA APP (no propone cargas):
  · cuota del histórico (las 90 de la carga del 07/10, con las filas de la revisión)
  · sueldo / anticipo de sueldo (sueldos.pagos, fecha + importe exactos)
  · anticipo a proveedor (fecha de pago + importe)
  · factura: por el CUIT de la leyenda; «importe exacto» o «mismo CUIT, otro importe» (información, no match)
Y una solapa «Resumen por CATEG» con fórmulas sobre la hoja: total, qué ya está, qué no, y el control.

    python scripts/trabajo_extracto_msa.py 1
Si existe la versión anterior, se traen las columnas «JMS:» que haya completado el usuario.
"""
import openpyxl, collections, datetime, json, re, sys, os, urllib.request, urllib.parse
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter as L

V = sys.argv[1] if len(sys.argv) > 1 else '1'
RAIZ = 'D:/Users/josem/Documents/Jose/Automatizarr/Claude/Control-Presupuestario-v1.1/'
DIR = RAIZ + '- Comunicacion JMS Claude - Archivos/Balance/Excels trabajando con claude/'
FUENTE = DIR + 'Extracto MSA - Control Presupuestario.xlsx'
SALIDA = DIR + f'Extracto MSA - Control Presupuestario - TRABAJO v{V}.xlsx'
PREVIA = DIR + f'Extracto MSA - Control Presupuestario - TRABAJO v{int(V) - 1}.xlsx'
VIEJA = RAIZ + '- Comunicacion JMS Claude - Archivos/Balance/extracto desde julio del 2025 hasta enero 26 parcial. galicia MSA cta cte pesos.xlsx'
TANDA1 = DIR + 'Tanda 1 - propuesta de carga v3.xlsx'
INICIO_EJ = datetime.datetime(2025, 7, 1)

env = {}
for l in open(RAIZ + '.env.local', encoding='utf-8'):
    if '=' in l and not l.strip().startswith('#'):
        k, v = l.split('=', 1); env[k.strip()] = v.strip().strip('"').strip("'")
def rest(tabla, params, schema=None):
    url = env['NEXT_PUBLIC_SUPABASE_URL'] + '/rest/v1/' + tabla + '?' + urllib.parse.urlencode(params, safe='(),.*:')
    h = {'apikey': env['SUPABASE_SERVICE_ROLE_KEY'], 'Authorization': 'Bearer ' + env['SUPABASE_SERVICE_ROLE_KEY']}
    if schema: h['Accept-Profile'] = schema
    return json.loads(urllib.request.urlopen(urllib.request.Request(url, headers=h)).read())

r2 = lambda x: round(float(x), 2) if isinstance(x, (int, float)) else 0.0
iso = lambda d: d.strftime('%Y-%m-%d') if isinstance(d, (datetime.datetime, datetime.date)) else str(d)[:10]
dmy = lambda s_: '/'.join(reversed(str(s_)[:10].split('-')))
s = lambda x: '' if x is None else str(x)
money = lambda x: f'{x:,.2f}'.replace(',', 'X').replace('.', ',').replace('X', '.')

# ── 1. La planilla del usuario (con fórmulas: se guarda tal cual) ─────────────────────────────────────────
wb = openpyxl.load_workbook(FUENTE)
ws = wb.worksheets[0]
wd = openpyxl.load_workbook(FUENTE, data_only=True).worksheets[0]
HR = 4
H = {s(c.value): c.column for c in ws[HR] if c.value}
movs = {}
for i in range(HR + 1, ws.max_row + 1):
    f = wd.cell(i, 1).value
    if not isinstance(f, datetime.datetime): continue
    g = lambda n: wd.cell(i, H[n]).value
    movs[i] = {'fila': i, 'fecha': f, 'deb': r2(g('Débitos')), 'cre': r2(g('Créditos')), 'saldo': r2(g('Saldo')),
               'desc': s(g('Descripción')), 'categ': s(g('CATEG')).strip(), 'det': s(g('Detalle')),
               'ley': ' '.join(s(g(k)) for k in ('Leyendas Adicionales1', 'Leyendas Adicionales2', 'Leyendas Adicionales3', 'Leyendas Adicionales4'))}
ultima_fila_mov = max(movs)
ej = lambda m: '25/26' if m['fecha'] >= INICIO_EJ else '24/25'
en_ej = [m for m in movs.values() if ej(m) == '25/26']

# ── 2. Las 90 cuotas del histórico → filas (vía la planilla vieja, que es la que usó la carga) ──────────────
wv = openpyxl.load_workbook(VIEJA, data_only=True).worksheets[0]
hv = {s(c.value): c.column for c in wv[3] if c.value}
vieja = {}
for i in range(4, wv.max_row + 1):
    f = wv.cell(i, 1).value
    if isinstance(f, datetime.datetime):
        vieja[i] = {'fecha': f, 'deb': r2(wv.cell(i, hv['Débitos']).value), 'cre': r2(wv.cell(i, hv['Créditos']).value),
                    'saldo': r2(wv.cell(i, hv['Saldo']).value), 'categ': s(wv.cell(i, hv['CATEG']).value).strip(),
                    'desc': s(wv.cell(i, hv['Descripción']).value), 'det': s(wv.cell(i, hv['Detalle']).value)}
clave = lambda m: (iso(m['fecha']), m['deb'], m['cre'], m['saldo'])
por_clave = collections.defaultdict(list)
for i, m in movs.items(): por_clave[clave(m)].append(i)
vieja_a_nueva = {}
for i, m in vieja.items():
    l = por_clave.get(clave(m))
    if not l: raise SystemExit(f'⚠ la fila {i} de la planilla vieja no está en la nueva: no sigo')
    vieja_a_nueva[i] = l.pop(0)

tpls = {t['id']: t for t in rest('egresos_sin_factura', {'select': 'id,nombre_referencia,responsable'})}
anteriores = rest('cuotas_egresos_sin_factura', {'select': 'id,egreso_id,fecha_estimada,monto,tipo_movimiento,descripcion,created_at', 'estado': 'eq.anterior'})
for c in anteriores:
    c['tpl'] = tpls.get(c['egreso_id'], {}).get('nombre_referencia', '?'); c['monto'] = r2(float(c['monto'] or 0)); c['filas'] = []
wt = openpyxl.load_workbook(TANDA1, data_only=True)['01 Cuotas propuestas']
ht = {c.value: i for i, c in enumerate(wt[1])}
usadas = set()
for row in wt.iter_rows(min_row=2, values_only=True):
    if not isinstance(row[0], int): continue
    monto = r2(row[ht['Débito (egreso)']] or row[ht['Crédito (ingreso)']] or 0)
    c = next((c for c in anteriores if c['id'] not in usadas and c['tpl'] == row[ht['Template']] and c['fecha_estimada'] == iso(row[ht['Fecha']])
              and c['tipo_movimiento'] == row[ht['Tipo']] and abs(c['monto'] - monto) < 0.01), None)
    if not c: raise SystemExit(f'⚠ cuota de la tanda 1 no encontrada en la base: {row[:6]}')
    usadas.add(c['id']); c['filas'] = [vieja_a_nueva[int(x)] for x in re.findall(r'\d+', s(row[ht['Filas de tu planilla']]))]; c['tanda'] = 'tanda 1'
for c in anteriores:
    if c['id'] in usadas: continue
    mes, ing = c['fecha_estimada'][:7], c['tipo_movimiento'] == 'ingreso'
    if 'FIMA' in c['tpl']:
        c['filas'] = [vieja_a_nueva[i] for i, m in vieja.items() if m['categ'].upper() == 'FCI' and iso(m['fecha'])[:7] == mes and ((m['cre'] > 0) if ing else (m['deb'] > 0))]; c['tanda'] = 'FCI'
    elif 'Plazo Fijo' in c['tpl']:
        c['filas'] = [vieja_a_nueva[i] for i, m in vieja.items() if re.search(r'plazo fij', m['desc'] + m['det'], re.I) and iso(m['fecha']) == c['fecha_estimada']
                      and abs((m['cre'] if ing else m['deb']) - c['monto']) < 0.01]; c['tanda'] = 'plazo fijo'
    else:
        c['tanda'] = 'ya estaba antes (creada ' + dmy(c['created_at']) + ')'
app = {}   # fila -> (tipo, detalle, referencia)
for c in anteriores:
    suma = sum((movs[f]['cre'] if c['tipo_movimiento'] == 'ingreso' else movs[f]['deb']) for f in c['filas'])
    if c['filas'] and abs(suma - c['monto']) > 0.01: raise SystemExit(f'⚠ la cuota {c["tpl"]} {c["fecha_estimada"]} no cierra con sus filas')
    for f in c['filas']:
        app[f] = ('Cuota del histórico', f'{c["tpl"]} · cuota {dmy(c["fecha_estimada"])} {c["tipo_movimiento"]} ${money(c["monto"])} ({c["tanda"]}, {len(c["filas"])} mov.)', c['id'])
cuotas_previas = [c for c in anteriores if not c['filas']]

# ── 3. Sueldos y anticipos de sueldo (fecha + importe exactos) ─────────────────────────────────────────────
emps = {e['id']: e['nombre'] for e in rest('empleados', {'select': 'id,nombre'}, 'sueldos')}
pagos = rest('pagos', {'select': 'id,empleado_id,fecha,monto,tipo,estado', 'and': '(fecha.gte.2025-07-01,fecha.lte.2026-01-31)'}, 'sueldos')
for p in pagos:
    f = next((i for i, m in movs.items() if i not in app and iso(m['fecha']) == p['fecha'] and abs(m['deb'] - r2(float(p['monto']))) < 0.01), None)
    if f: app[f] = ('Sueldo', f'{emps.get(p["empleado_id"], "?")} · {p["tipo"]} {dmy(p["fecha"])} ${money(r2(float(p["monto"])))} ({p["estado"]})', p['id'])
    else: print('⚠ pago de sueldo sin movimiento:', emps.get(p['empleado_id']), p['fecha'], p['monto'])

# ── 4. Anticipos a proveedores ────────────────────────────────────────────────────────────────────────────
ants = rest('anticipos_proveedores', {'select': 'id,nombre_proveedor,cuit_proveedor,monto,monto_sicore,fecha_pago,estado', 'and': '(fecha_pago.gte.2025-07-01,fecha_pago.lte.2026-01-31)'})
for a in ants:
    netos = {r2(float(a['monto'] or 0)), r2(float(a['monto'] or 0) - float(a['monto_sicore'] or 0))}
    f = next((i for i, m in movs.items() if i not in app and iso(m['fecha']) == s(a['fecha_pago'])[:10] and any(abs(m['deb'] - n) < 0.01 for n in netos)), None)
    if f: app[f] = ('Anticipo a proveedor', f'{a["nombre_proveedor"]} · {dmy(a["fecha_pago"])} ${money(r2(float(a["monto"])))} ({a["estado"]})', a['id'])

# ── 5. Facturas: por el CUIT de la leyenda ────────────────────────────────────────────────────────────────
fcs = rest('comprobantes_arca', {'select': 'id,cuit,denominacion_emisor,fecha_emision,tipo_comprobante,punto_venta,numero_desde,imp_total,monto_sicore,estado',
                                  'and': '(fecha_emision.gte.2024-10-01,fecha_emision.lte.2026-01-31)'}, 'msa')
por_cuit = collections.defaultdict(list)
for x in fcs: por_cuit[re.sub(r'\D', '', s(x['cuit']))].append(x)
usadas_fc = set(); info_cuit = {}
def nro(x): return f'{x["tipo_comprobante"]}-{int(x["punto_venta"] or 0):04d}-{int(x["numero_desde"] or 0):08d}'
for i, m in sorted(movs.items(), key=lambda x: x[1]['fecha']):
    if ej(m) != '25/26' or i in app or m['deb'] <= 0: continue
    cuits = [c for c in re.findall(r'\b\d{11}\b', m['ley']) if c in por_cuit]
    if not cuits: continue
    cands = [x for c in cuits for x in por_cuit[c] if x['id'] not in usadas_fc and s(x['fecha_emision']) <= iso(m['fecha'] + datetime.timedelta(days=5))]
    ex = next((x for x in cands if any(abs(m['deb'] - v) < 0.01 for v in (r2(float(x['imp_total'])), r2(float(x['imp_total']) - float(x['monto_sicore'] or 0))))), None)
    if ex:
        usadas_fc.add(ex['id'])
        app[i] = ('Factura (importe exacto)', f'{ex["denominacion_emisor"]} · {nro(ex)} del {dmy(ex["fecha_emision"])} ${money(r2(float(ex["imp_total"])))} ({ex["estado"]})', ex['id'])
    else:
        todas = [x for c in cuits for x in por_cuit[c]]
        cerca = sorted(todas, key=lambda x: abs(r2(float(x['imp_total'])) - m['deb']))[:3]
        info_cuit[i] = ('Proveedor con facturas en la app (otro importe)',
                        f'{todas[0]["denominacion_emisor"]}: {len(todas)} factura(s) entre oct-24 y ene-26; las más cercanas: ' +
                        '; '.join(f'{nro(x)} {dmy(x["fecha_emision"])} ${money(r2(float(x["imp_total"])))} ({x["estado"]})' for x in cerca), '')

# ── 6. Columnas nuevas a la derecha (fila 4 = encabezados; sólo filas de movimientos) ──────────────────────
COL0 = ws.max_column + 2               # después de TODO lo que usa el usuario (tiene marcas sueltas a la derecha), con una columna libre
NUEVAS = ['Ejercicio', 'App: ¿ya está?', 'App: qué es', 'App: id', 'Propuesta: destino', 'Propuesta: cuenta', 'Propuesta: cómo se identificó', 'Aviso', 'JMS: decisión', 'JMS: nota']
C = {n: COL0 + k for k, n in enumerate(NUEVAS)}
AZUL = PatternFill('solid', fgColor='DDEBF7'); VERDE = PatternFill('solid', fgColor='E2EFDA'); AMAR = PatternFill('solid', fgColor='FFF2CC'); GRIS = PatternFill('solid', fgColor='EDEDED'); B = Font(bold=True)
fill_de = lambda n: AZUL if n.startswith('App') or n == 'Ejercicio' else (VERDE if n.startswith('Propuesta') else (AMAR if n.startswith('JMS') else GRIS))
for n, c in C.items():
    cell = ws.cell(HR, c, n); cell.font = B; cell.fill = fill_de(n)
# lo que el usuario haya escrito en la versión anterior
previas = {}
if os.path.exists(PREVIA):
    wp = openpyxl.load_workbook(PREVIA, data_only=True).worksheets[0]
    hp = {s(c.value): c.column for c in wp[HR] if c.value}
    for i in range(HR + 1, wp.max_row + 1):
        f = wp.cell(i, 1).value
        if not isinstance(f, datetime.datetime): continue
        k = (iso(f), r2(wp.cell(i, hp['Débitos']).value), r2(wp.cell(i, hp['Créditos']).value), r2(wp.cell(i, hp['Saldo']).value))
        previas[k] = {n: wp.cell(i, hp[n]).value for n in NUEVAS if n.startswith('JMS') and n in hp}
cuenta = collections.Counter()
for i, m in movs.items():
    e = ej(m); ws.cell(i, C['Ejercicio'], e)
    if e == '24/25':
        ws.cell(i, C['App: ¿ya está?'], '— ejercicio anterior (no se carga)')
    elif i in app:
        t, d, ref = app[i]; ws.cell(i, C['App: ¿ya está?'], t); ws.cell(i, C['App: qué es'], d); ws.cell(i, C['App: id'], ref)
    elif i in info_cuit:
        t, d, _ = info_cuit[i]; ws.cell(i, C['App: ¿ya está?'], t); ws.cell(i, C['App: qué es'], d)
    else:
        ws.cell(i, C['App: ¿ya está?'], 'No está en la app')
    cuenta[ws.cell(i, C['App: ¿ya está?']).value] += 1
    for n, v in previas.get(clave(m), {}).items():
        if v not in (None, ''): ws.cell(i, C[n], v)
ws.auto_filter.ref = f'A{HR}:{L(C["JMS: nota"])}{ultima_fila_mov}'

# ── 7. Resumen por CATEG (fórmulas sobre la hoja) ─────────────────────────────────────────────────────────
rs = wb.create_sheet('Resumen por CATEG')
N = ws.title.replace("'", "''")
rng = lambda col: f"'{N}'!${L(col)}${HR + 1}:${L(col)}${ultima_fila_mov}"
cCAT, cDEB, cCRE, cEJ, cAPP = H['CATEG'], H['Débitos'], H['Créditos'], C['Ejercicio'], C['App: ¿ya está?']
TIPOS = ['Cuota del histórico', 'Sueldo', 'Anticipo a proveedor', 'Factura (importe exacto)', 'Proveedor con facturas en la app (otro importe)', 'No está en la app']
rs.append([f'RESUMEN POR CATEG — ejercicio 25/26 (jul-2025 → ene-2026) de «{ws.title}». Todo con fórmulas sobre la hoja: si cambiás una CATEG allá, cambia acá. Generado {datetime.date.today():%d/%m/%Y}. v{V}: sólo marca lo que YA está en la app; no propone cargas.'])
rs.append(['«Factura (importe exacto)» = un comprobante de ese CUIT por el mismo importe (o el importe menos la retención). «Proveedor con facturas (otro importe)» es sólo información: el CUIT tiene facturas en la app pero ninguna por ese importe — hay que mirarlo.'])
rs.append([])
enc = ['CATEG', 'Movimientos', 'Débitos', 'Créditos']
for t in TIPOS: enc += [f'{t} · cant.', f'{t} · débitos', f'{t} · créditos']
enc += ['Control: movimientos − suma de las partes (tiene que dar 0)']
rs.append(enc); hr = rs.max_row
for cell in rs[hr]: cell.font = B; cell.fill = AZUL
cats = sorted({m['categ'] for m in en_ej}, key=lambda x: x.upper())
ini = hr + 1
for cat in cats:
    n = rs.max_row + 1
    fila = [cat if cat else '(sin CATEG)',
            f'=COUNTIFS({rng(cCAT)},{"$A" + str(n) if cat else chr(34) + chr(34)},{rng(cEJ)},"25/26")',
            f'=SUMIFS({rng(cDEB)},{rng(cCAT)},{"$A" + str(n) if cat else chr(34) + chr(34)},{rng(cEJ)},"25/26")',
            f'=SUMIFS({rng(cCRE)},{rng(cCAT)},{"$A" + str(n) if cat else chr(34) + chr(34)},{rng(cEJ)},"25/26")']
    crit = f'{rng(cCAT)},{"$A" + str(n) if cat else chr(34) + chr(34)},{rng(cEJ)},"25/26"'
    for t in TIPOS:
        fila += [f'=COUNTIFS({crit},{rng(cAPP)},"{t}")', f'=SUMIFS({rng(cDEB)},{crit},{rng(cAPP)},"{t}")', f'=SUMIFS({rng(cCRE)},{crit},{rng(cAPP)},"{t}")']
    partes = '-'.join(f'{L(5 + 3 * k)}{n}' for k in range(len(TIPOS)))
    fila.append(f'=B{n}-{partes}')
    rs.append(fila)
fin = rs.max_row
rs.append(['TOTAL'] + [f'=SUM({L(j)}{ini}:{L(j)}{fin})' for j in range(2, len(enc) + 1)])
tot = rs.max_row
rs.append([])
rs.append(['CONTROL contra la hoja', 'movimientos 25/26 en la hoja', 'débitos', 'créditos'])
rs.append(['', f'=COUNTIFS({rng(cEJ)},"25/26")', f'=SUMIFS({rng(cDEB)},{rng(cEJ)},"25/26")', f'=SUMIFS({rng(cCRE)},{rng(cEJ)},"25/26")'])
rs.append(['diferencia con el TOTAL (tiene que dar 0)', f'=B{tot}-B{tot + 3}', f'=ROUND(C{tot}-C{tot + 3},2)', f'=ROUND(D{tot}-D{tot + 3},2)'])
rs.append([])
rs.append(['Cuotas «anterior» que YA ESTABAN antes de la carga del histórico y no salen de este extracto (no duplicar):'])
for c in cuotas_previas:
    rs.append(['', c['tpl'], dmy(c['fecha_estimada']), c['monto'], c['tanda'], c['id']])
for row in rs.iter_rows(min_row=ini, max_row=rs.max_row):
    for cell in row:
        if cell.column >= 3 and isinstance(cell.value, str) and cell.value.startswith('=SUMIFS'): cell.number_format = '#,##0.00'
for r in (tot + 3, tot + 4):
    for j in (3, 4): rs.cell(r, j).number_format = '#,##0.00'
rs.freeze_panes = f'B{hr + 1}'
rs.column_dimensions['A'].width = 22
wb.save(SALIDA)

print(SALIDA)
print('movimientos', len(movs), '| 25/26', len(en_ej), '| filas con cuota del histórico', sum(1 for v in app.values() if v[0] == 'Cuota del histórico'))
for k, v in cuenta.most_common(): print(f'  {v:5} {k}')
print('cuotas previas (no salen del extracto):', [(c['tpl'], c['fecha_estimada'], c['monto']) for c in cuotas_previas])
